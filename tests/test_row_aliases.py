"""Optional source-row aliases, fallback, precedence and v3 migration."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.models import SourceField
from excel_copy_script.workbook import WorkbookSession


class RowSpecificAliasTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.folder = Path(tmp.name)
        self.destination = self.folder / "summary.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        ws["A3"] = "№"
        ws["B3"] = "Резерв"
        ws["C3"] = "Техніка"
        ws["D3"] = "Додатково"
        wb.save(self.destination)
        self.alias_file = self.folder / "aliases.json"
        self.aliases = AliasStore(self.alias_file)
        self.session = WorkbookSession(self.destination, self.aliases)

    def input_file(self, mapping, sheet="ВИСНОВКИ", filename="input.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for cell, value in mapping.items():
            ws[cell] = value
        path = self.folder / filename
        wb.save(path)
        return path

    def set_alias(self, name="Повтор", source_column=3, target=2, row=None, sheet="ВИСНОВКИ"):
        source = SourceField(row or 2, name, 1, sheet=sheet, value_column=source_column)
        destination = next(f for f in self.session.fields if f.column == target)
        self.aliases.set(self.session.template_id, source, destination, row=row)

    def test_general_and_scoped_aliases_can_coexist_and_survive_reload(self):
        self.set_alias()
        self.set_alias(row=14, source_column=4)
        self.set_alias(row=19, target=3, source_column=5)
        key = self.session.template_id
        self.assertEqual(self.aliases.get(key, "Повтор")["source_column"], 3)
        self.assertEqual(self.aliases.get(key, "Повтор", row=14)["source_column"], 4)
        self.assertEqual(self.aliases.get(key, "Повтор", row=19)["source_column"], 5)
        self.assertEqual(self.aliases.get(key, "Повтор", row=15)["source_column"], 3)
        self.assertIsNone(self.aliases.get_exact(key, "Повтор", row=15))
        loaded = AliasStore(self.alias_file)
        self.assertEqual([(entry.source_row, entry.source_column) for entry
                          in loaded.list_for_template(key)], [(None, 3), (14, 4), (19, 5)])

    def test_scoped_alias_overrides_general_for_exact_row_only(self):
        self.set_alias(target=2)
        self.set_alias(row=14, target=3, source_column=4)
        file = self.input_file({
            "B14": "Повтор", "C14": 11, "D14": 99,
            "B15": "Повтор", "C15": 55, "D15": 77,
        })
        result = self.session.open_input(file)
        self.assertEqual({
            r.destination.column: (r.source.row, r.source.value) for r in result.copied
        }, {2: (15, 55), 3: (14, 99)})
        self.assertEqual(self.session.sheet["B4"].value, 55)
        self.assertEqual(self.session.sheet["C4"].value, 99)
        self.assertFalse(result.alias_issues)

    def test_same_destination_one_specific_beats_general_duplicate(self):
        self.set_alias(target=3)
        self.set_alias(row=14, target=3, source_column=4)
        file = self.input_file({
            "B14": "Повтор", "C14": 11, "D14": 99,
            "B19": "Повтор", "C19": 55,
        })
        result = self.session.open_input(file)
        self.assertEqual([(r.source.row, r.source.value) for r in result.copied],
                         [(14, 99)])
        self.assertEqual(self.session.sheet["C4"].value, 99)
        self.assertIn("Пріоритет", result.alias_issues[("висновки", 19)])
        self.assertNotIn(("висновки", 14), result.alias_issues)

    def test_specific_alias_disambiguates_even_if_two_other_general_rows_match(self):
        self.set_alias(target=3)
        self.set_alias(row=14, target=3)
        result = self.session.open_input(self.input_file({
            "B14": "Повтор", "C14": 10,
            "B19": "Повтор", "C19": 20,
            "B23": "Повтор", "C23": 30,
        }))
        self.assertEqual([r.source.row for r in result.copied], [14])
        self.assertIn(("висновки", 19), result.alias_issues)
        self.assertIn(("висновки", 23), result.alias_issues)

    def test_different_names_do_not_get_row_precedence(self):
        self.set_alias(target=3)
        self.set_alias(name="Інший", row=14, target=3)
        result = self.session.open_input(self.input_file({
            "B14": "Інший", "C14": 11,
            "B19": "Повтор", "C19": 55,
        }))
        self.assertFalse(result.copied)
        self.assertEqual(len(result.alias_issues), 2)

    def test_two_row_specific_candidates_same_destination_still_conflict(self):
        self.set_alias(row=14, target=3)
        self.set_alias(row=19, target=3)
        result = self.session.open_input(self.input_file({
            "B14": "Повтор", "C14": 11,
            "B19": "Повтор", "C19": 55,
        }))
        self.assertFalse(result.copied)
        self.assertEqual(len(result.alias_issues), 2)

    def test_wrong_row_or_wrong_name_or_sheet_does_not_match_scoped(self):
        self.set_alias(row=14, target=3)
        for index, (sheet, row, name) in enumerate((
            ("ВИСНОВКИ", 13, "Повтор"),
            ("ВИСНОВКИ", 14, "Інша назва"),
            ("Склад", 14, "Повтор"),
        )):
            result = self.session.open_input(self.input_file(
                {f"B{row}": name, f"C{row}": 10},
                sheet=sheet, filename=f"mismatch{index}.xlsx",
            ))
            self.assertFalse(result.copied)
            self.assertEqual(self.session.source_alias_targets(), {})

    def test_delete_exact_variant_preserves_other_variant(self):
        self.set_alias()
        self.set_alias(row=14)
        key = self.session.template_id
        self.assertTrue(self.aliases.delete(key, "Повтор", row=14))
        self.assertIsNone(self.aliases.get_exact(key, "Повтор", row=14))
        self.assertIsNotNone(self.aliases.get(key, "Повтор", row=14))
        self.set_alias(row=14)
        self.assertTrue(self.aliases.delete(key, "Повтор"))
        self.assertIsNone(self.aliases.get(key, "Повтор"))
        self.assertIsNotNone(self.aliases.get(key, "Повтор", row=14))
        self.assertEqual(len(AliasStore(self.alias_file).list_for_template(key)), 1)

    def test_invalid_scoped_row_is_rejected(self):
        for row in (0, -1, 1.5, True):
            with self.assertRaises(ValueError):
                self.set_alias(row=row)
        self.assertFalse(self.aliases.list_for_template(self.session.template_id))

    def test_manual_copy_creates_scoped_alias_and_keeps_general(self):
        self.set_alias(target=2)
        file = self.input_file({"B14": "Повтор", "C14": 11})
        self.session.open_input(file)
        # Existing alias already copied. Instead use a fresh row with a
        # different name, then add a scoped alias for it.
        new = self.input_file({"B14": "Особовий", "F14": 42}, filename="manual.xlsx")
        self.session.open_input(new)
        self.session.copy_manually(14, 3, source_column=6, alias_row=14)
        record = self.session.current.copied[0]
        self.assertEqual(record.alias_row, 14)
        self.assertEqual(self.aliases.get_exact(
            self.session.template_id, "Особовий", row=14,
        )["source_column"], 6)
        self.assertIsNone(self.aliases.get_exact(
            self.session.template_id, "Особовий",
        ))
        result = self.session.open_input(new)
        self.assertEqual([(r.source.row, r.destination.column) for r in result.copied],
                         [(14, 3)])

    def test_v3_migrates_as_unrestricted_and_preserves_backup(self):
        old = {"version": 3, "templates": {
            self.session.template_id: {"висновки": {
                "повтор": {"header": "техніка", "column": 3,
                            "sheet_name": "ВИСНОВКИ", "source_name": "Повтор",
                            "source_column": 5}
            }}
        }}
        original = json.dumps(old, ensure_ascii=False)
        self.alias_file.write_text(original, encoding="utf-8")
        migrated = AliasStore(self.alias_file)
        self.assertEqual(migrated.data["version"], 4)
        self.assertEqual(migrated.get(self.session.template_id, "Повтор", row=150)["source_column"], 5)
        self.assertIsNone(migrated.get_exact(self.session.template_id, "Повтор", row=150))
        self.assertEqual(migrated.list_for_template(self.session.template_id)[0].source_row, None)
        self.assertEqual(
            (self.folder / "aliases.json.v3.bak").read_text(encoding="utf-8"),
            original,
        )

    def test_alias_save_failure_rolls_back_scoped_set_and_delete(self):
        self.set_alias()
        key = self.session.template_id
        with patch.object(self.aliases, "_save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                self.set_alias(row=14)
            with self.assertRaises(OSError):
                self.aliases.delete(key, "Повтор")
        self.assertIsNone(self.aliases.get_exact(key, "Повтор", row=14))
        self.assertIsNotNone(self.aliases.get(key, "Повтор"))


if __name__ == "__main__":
    unittest.main()
