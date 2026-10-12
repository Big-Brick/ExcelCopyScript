"""Tk profile switching, persisted last profile and layout update."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from excel_copy_script.gui import ExcelCopyApp
from excel_copy_script.profiles import Profile, ProfileManager
from excel_copy_script.workbook_layout import WorkbookLayout


class ProfileGuiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.manager = ProfileManager(self.root / "config")
        self.patch = patch(
            "excel_copy_script.gui.ProfileManager",
            return_value=self.manager,
        )
        self.patch.start()
        self.addCleanup(self.patch.stop)
        self.app = ExcelCopyApp()
        self.addCleanup(self.app.destroy)
        self.app.geometry("1400x860")
        self.app.update()

    def test_last_profile_is_loaded_at_launch(self):
        self.assertEqual(self.app.profile.path, self.manager.default_path.resolve())
        self.assertTrue(self.app.profile.path.exists())
        self.assertEqual(
            self.app.create_sources.value_combo.cget("values"),
            ("C", "D", "E", "F", "G"),
        )

    def test_open_profile_changes_columns_and_remembers_choice(self):
        alt = self.app.profile.copy_as(self.root / "alternate.json")
        alt.layout = WorkbookLayout(
            destination_sheet="Report", header_row=5, first_data_row=6,
            first_destination_column=3, source_name_column=4,
            first_source_row=2, source_value_first_column=6,
            source_value_last_column=8, preferred_source_sheet="Data",
        )
        alt.save()
        with patch("excel_copy_script.gui.filedialog.askopenfilename",
                   return_value=str(alt.path)):
            self.app.open_profile()
        self.app.update()
        self.assertEqual(self.app.profile.path, alt.path)
        self.assertEqual(self.manager.load_last().path, alt.path)
        self.assertEqual(self.app.create_sources.COLUMNS, (6, 7, 8))
        self.assertEqual(tuple(self.app.create_sources.values_tree["columns"]),
                         ("F", "G", "H"))
        self.assertEqual(self.app.create_sources.selected_value_column.get(), "F")
        self.app.create_sources.set_sheets(["Other", "Data"])
        self.assertEqual(self.app.create_sources.selected_sheet.get(), "Data")
        self.assertIn("alternate.json", self.app.profile_var.get())

    def test_invalid_profile_does_not_switch_or_replace_active_file(self):
        malformed = self.root / "bad.json"
        malformed.write_text('{"version": 4, "templates": {}}')
        old_path = self.app.profile.path
        with (
            patch("excel_copy_script.gui.filedialog.askopenfilename",
                  return_value=str(malformed)),
            patch("excel_copy_script.gui.messagebox.showerror") as error,
        ):
            self.app.open_profile()
        error.assert_called_once()
        self.assertEqual(self.app.profile.path, old_path)

    def test_switch_with_open_destination_never_writes_excel(self):
        dest = self.root / "sheet.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["B4"] = "№", "Обсяг", 999
        wb.save(dest)
        with patch("excel_copy_script.gui.filedialog.askopenfilename",
                   return_value=str(dest)):
            self.app.open_destination()
        original = dest.read_bytes()
        first = self.app.profile
        alternative = first.copy_as(self.root / "compatible.json")
        self.app._activate_profile(alternative)
        self.assertEqual(dest.read_bytes(), original)
        self.assertIsNotNone(self.app.session)
        self.assertIsNone(self.app.session.current)
        self.assertEqual(self.app.opened_sources, [])

        alternative = first.copy_as(self.root / "incompatible.json")
        alternative.layout = WorkbookLayout(destination_sheet="NotHere")
        alternative.save()
        self.app._activate_profile(alternative)
        self.assertEqual(dest.read_bytes(), original)
        self.assertIsNone(self.app.session)
        self.assertEqual(str(self.app.source_button.cget("state")), "disabled")
        self.assertTrue(self.app.profile_var.get().endswith("incompatible.json"))


if __name__ == "__main__":
    unittest.main()
