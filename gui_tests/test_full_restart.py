"""Xvfb regression: full restart button replays files without file dialogs."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.gui import ExcelCopyApp
from excel_copy_script.models import SourceField
from excel_copy_script.recent_directories import RecentDirectories


class FullRestartGuiTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.folder = Path(tmp.name)
        self.summary = self.folder / "summary.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"], ws["D3"] = (
            "№", "Особовий склад", "Техніка", "Резерв",
        )
        wb.save(self.summary)
        self.source_a = self.folder / "a.xlsx"
        self.source_b = self.folder / "b.xlsx"
        self.create_source(self.source_a, "ВИСНОВКИ", [
            (1, "Чисельність", 11), (2, "Запас", 100),
        ])
        self.create_source(self.source_b, "Склад", [(1, "Бронетехніка", 22)])
        self.aliases = AliasStore(self.folder / "aliases.json")
        self.app = ExcelCopyApp()
        self.addCleanup(self.app.destroy)
        self.app.recent_directories = RecentDirectories(self.folder / "recent.json")
        self.app.geometry("1400x800")
        self.app.update()

    @staticmethod
    def create_source(path, sheet, rows):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for row, name, value in rows:
            ws.cell(row, 2).value = name
            ws.cell(row, 3).value = value
        wb.save(path)

    def load_files(self):
        with (
            patch("excel_copy_script.gui.AliasStore", return_value=self.aliases),
            patch("excel_copy_script.gui.filedialog.askopenfilename",
                  side_effect=[str(self.summary), str(self.source_a), str(self.source_b)]),
        ):
            self.app.open_destination()
            self.aliases.set(
                self.app.session.template_id,
                SourceField(1, "Чисельність", 11),
                self.app.session.fields[0],
            )
            self.aliases.set(
                self.app.session.template_id,
                SourceField(1, "Бронетехніка", 22, sheet="Склад"),
                self.app.session.fields[1],
            )
            self.app.open_source()
            self.app.open_source()
        self.app.update()

    def test_button_replays_in_order_with_latest_aliases_and_no_dialog(self):
        self.assertEqual(str(self.app.restart_button["state"]), "disabled")
        self.load_files()
        self.assertEqual(str(self.app.restart_button["state"]), "normal")
        self.assertEqual(self.app.opened_sources, [self.source_a, self.source_b])

        # Changing aliases must affect the new run, not just GUI highlights.
        self.aliases.delete(self.app.session.template_id, "Чисельність")
        self.aliases.set(
            self.app.session.template_id,
            SourceField(2, "Запас", 100),
            self.app.session.fields[2],
        )
        self.app.notebook.select(1)
        self.app.update()
        previous_session = self.app.session
        with patch(
            "excel_copy_script.gui.filedialog.askopenfilename",
            side_effect=AssertionError("Restart must not open file dialogs"),
        ):
            self.app.restart_button.invoke()
        self.app.update()

        self.assertIsNot(self.app.session, previous_session)
        self.assertEqual(self.app.notebook.index("current"), 0)
        self.assertEqual(self.app.opened_sources, [self.source_a, self.source_b])
        self.assertEqual(self.app.session.current.filename, "b.xlsx")
        self.assertEqual(self.app.session.current.row, 5)
        self.assertIn("повторно імпортовано 2", self.app.status_var.get())
        ws = load_workbook(self.summary)["Дані_підрозділів"]
        self.assertIsNone(ws["B4"].value)
        self.assertEqual(ws["D4"].value, 100)
        self.assertEqual(ws["C5"].value, 22)
        self.assertIn("b.xlsx", self.app.source_var.get())

    def test_failed_restart_keeps_workbook_ui_and_file_history(self):
        self.load_files()
        previous_session = self.app.session
        previous_status = self.app.status_var.get()
        previous_bytes = self.summary.read_bytes()
        self.source_b.unlink()
        with (
            patch("excel_copy_script.gui.messagebox.showerror") as error,
            patch(
                "excel_copy_script.gui.filedialog.askopenfilename",
                side_effect=AssertionError("Unexpected dialog"),
            ),
        ):
            self.app.restart_button.invoke()
        self.app.update()
        error.assert_called_once()
        self.assertIn("b.xlsx", str(error.call_args))
        self.assertIs(self.app.session, previous_session)
        self.assertEqual(self.app.status_var.get(), previous_status)
        self.assertEqual(self.app.opened_sources, [self.source_a, self.source_b])
        self.assertEqual(self.summary.read_bytes(), previous_bytes)
        self.assertFalse(list(self.folder.glob(".excel-copy-restart-*.xlsx")))

    def test_opening_new_destination_resets_replay_history(self):
        self.load_files()
        self.assertEqual(len(self.app.opened_sources), 2)
        with patch(
            "excel_copy_script.gui.filedialog.askopenfilename",
            return_value=str(self.summary),
        ), patch("excel_copy_script.gui.AliasStore", return_value=self.aliases):
            self.app.open_destination()
        self.assertEqual(self.app.opened_sources, [])
        self.assertEqual(str(self.app.restart_button["state"]), "normal")
        self.assertIsNone(self.app.session.current)
        self.app.restart_button.invoke()
        self.assertIsNone(self.app.session.current)
        self.assertIn("немає", self.app.status_var.get())
        self.assertFalse(list(self.folder.glob(".excel-copy-restart-*.xlsx")))


if __name__ == "__main__":
    unittest.main()
