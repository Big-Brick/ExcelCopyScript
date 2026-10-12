"""Worksheet-scoped source preview with frozen identifiers and scrollable values."""

from __future__ import annotations

import tkinter as tk
from collections.abc import Callable, Sequence
from tkinter import ttk

from openpyxl.utils import get_column_letter

from .config import SOURCE_SHEET
from .models import SourceField
from .utils import present


class SourceRowsPanel(ttk.Frame):
    """Preview all B:G source rows while keeping alias, row and name frozen."""

    COLUMNS = tuple(range(3, 8))

    def __init__(
        self, parent: tk.Misc, on_selection_changed: Callable[[], None] | None = None
    ) -> None:
        super().__init__(parent, padding=4)
        self._rows: list[SourceField] = []
        self._alias_targets: dict[tuple[str, int], str] = {}
        self._copied_locations: set[tuple[str, int]] = set()
        self._alias_issues: dict[tuple[str, int], str] = {}
        self._sources_by_item_id: dict[str, SourceField] = {}
        self._on_selection_changed = on_selection_changed
        self._syncing_selection = False
        self._syncing_scroll = False
        self.selected_sheet = tk.StringVar(value="")
        self.selected_value_column = tk.StringVar(value="C")

        ttk.Label(self, text="Вхідні рядки (усі)",
                  font=("", 10, "bold")).pack(anchor="w", pady=(0, 5))
        selector = ttk.Frame(self)
        selector.pack(fill="x", pady=(0, 6))
        ttk.Label(selector, text="Аркуш:").pack(side="left")
        self.sheet_combo = ttk.Combobox(
            selector, textvariable=self.selected_sheet, state="disabled", width=20,
        )
        self.sheet_combo.pack(side="left", fill="x", expand=True, padx=(6, 8))
        self.sheet_combo.bind("<<ComboboxSelected>>", self._on_sheet_changed)
        ttk.Label(selector, text="Значення з:").pack(side="left")
        self.value_combo = ttk.Combobox(
            selector, textvariable=self.selected_value_column,
            state="readonly", values=("C", "D", "E", "F", "G"), width=3,
        )
        self.value_combo.pack(side="left", padx=(6, 0))
        self.value_combo.bind("<<ComboboxSelected>>", self._notify_selection)

        container = ttk.Frame(self)
        container.pack(fill="both", expand=True)

        # A real draggable sash separates the frozen row identifiers from
        # the horizontally scrollable C:G values. Its position is independent
        # in each SourceRowsPanel (one panel per notebook tab).
        self.divider = ttk.Panedwindow(container, orient="horizontal")
        self.divider.pack(side="left", fill="both", expand=True)

        fixed = ttk.Frame(self.divider, width=395)
        fixed.pack_propagate(False)
        self.tree = ttk.Treeview(
            fixed, columns=("alias", "row", "name"),
            show="headings", selectmode="browse",
        )
        for key, title, width in (
            ("alias", "Куди (аліас)", 150),
            ("row", "Рядок", 55),
            ("name", "Назва", 190),
        ):
            self.tree.heading(key, text=title)
            self.tree.column(key, width=width, minwidth=50, stretch=False)
        self.tree.pack(fill="both", expand=True)
        # Reserve exactly the same height as the right-hand horizontal scrollbar,
        # so both Treeviews have identical visible row heights and remain aligned.
        frozen_spacer = ttk.Frame(fixed)
        frozen_spacer.pack(fill="x")

        scrollable = ttk.Frame(self.divider)
        letters = tuple(get_column_letter(col) for col in self.COLUMNS)
        self.values_tree = ttk.Treeview(
            scrollable, columns=letters, show="headings", selectmode="browse",
        )
        for letter in letters:
            self.values_tree.heading(letter, text=letter)
            self.values_tree.column(letter, width=140, minwidth=100, stretch=False)
        xscroll = ttk.Scrollbar(
            scrollable, orient="horizontal", command=self.values_tree.xview,
        )
        self.values_tree.configure(xscrollcommand=xscroll.set)
        self.values_tree.pack(fill="both", expand=True)
        xscroll.pack(fill="x")
        frozen_spacer.configure(height=xscroll.winfo_reqheight())

        # Keep the left pane's preferred initial width while allowing the
        # sash to be moved freely. The right pane consumes window resizes.
        self.divider.add(fixed, weight=0)
        self.divider.add(scrollable, weight=1)

        vertical = ttk.Scrollbar(container, orient="vertical", command=self._yview)
        vertical.pack(side="right", fill="y")
        self.vertical_scroll = vertical
        self.tree.configure(yscrollcommand=lambda first, last: self._sync_y(first, last, self.values_tree))
        self.values_tree.configure(yscrollcommand=lambda first, last: self._sync_y(first, last, self.tree))

        self.tree.bind("<<TreeviewSelect>>", lambda event: self._sync_selection(self.tree, self.values_tree))
        self.values_tree.bind("<<TreeviewSelect>>", lambda event: self._sync_selection(self.values_tree, self.tree))
        self.values_tree.bind("<ButtonRelease-1>", self._on_value_click)

        for table in (self.tree, self.values_tree):
            table.tag_configure("unusable", background="#e5e5e5", foreground="#777777")
            table.tag_configure("aliased", background="#ddf2df")
            table.tag_configure("skipped", background="#fff0c2")
        self.issue_var = tk.StringVar(value="")
        ttk.Label(
            self, textvariable=self.issue_var,
            foreground="#875800", wraplength=690,
        ).pack(fill="x", pady=(4, 0))

    def _yview(self, *args) -> None:
        self.tree.yview(*args)
        self.values_tree.yview(*args)

    def _sync_y(self, first: str, last: str, other: ttk.Treeview) -> None:
        self.vertical_scroll.set(first, last)
        if self._syncing_scroll:
            return
        self._syncing_scroll = True
        try:
            # Both Treeviews call this callback. Avoid resetting the peer's
            # viewport to the same position: redundant updates may recurse.
            position = float(first)
            if abs(other.yview()[0] - position) > 1e-6:
                other.yview_moveto(position)
        finally:
            self._syncing_scroll = False

    def _sync_selection(self, current: ttk.Treeview, other: ttk.Treeview) -> None:
        if self._syncing_selection:
            return
        self._syncing_selection = True
        try:
            selected = current.selection()
            # Tk's <<TreeviewSelect>> is queued and delivered *after* the
            # handler returns. A synchronous reentrancy flag alone therefore
            # cannot prevent an endless selection event ping-pong.
            #
            # selection_set() emits another event even when the selection
            # is unchanged, so touch the peer only if the two selections differ.
            if other.selection() != selected:
                if selected:
                    other.selection_set(selected[0])
                else:
                    other.selection_remove(*other.selection())
        finally:
            self._syncing_selection = False
        self._notify_selection()

    def _on_value_click(self, event) -> None:
        if self.values_tree.identify_region(event.x, event.y) != "cell":
            return
        index = int(self.values_tree.identify_column(event.x).lstrip("#")) - 1
        if 0 <= index < len(self.COLUMNS):
            self.selected_value_column.set(get_column_letter(self.COLUMNS[index]))
            self._notify_selection()

    def source_column(self) -> int:
        return ord(self.selected_value_column.get()) - ord("A") + 1

    def set_sheets(self, sheets: Sequence[str]) -> None:
        sheets = list(sheets)
        self.sheet_combo.configure(values=sheets, state="readonly" if sheets else "disabled")
        self.selected_sheet.set(
            SOURCE_SHEET if SOURCE_SHEET in sheets else (sheets[0] if sheets else "")
        )
        self.selected_value_column.set("C")
        self.refresh(
            [], alias_targets={}, copied_locations=set(), alias_issues={},
        )

    def refresh(
        self, rows: Sequence[SourceField] | None = None,
        alias_targets: dict[tuple[str, int], str] | None = None,
        copied_locations: set[tuple[str, int]] | None = None,
        alias_issues: dict[tuple[str, int], str] | None = None,
    ) -> None:
        if rows is not None:
            self._rows = list(rows)
        if alias_targets is not None:
            self._alias_targets = dict(alias_targets)
        if copied_locations is not None:
            self._copied_locations = set(copied_locations)
        if alias_issues is not None:
            self._alias_issues = dict(alias_issues)
        self._sources_by_item_id.clear()
        self.tree.delete(*self.tree.get_children())
        self.values_tree.delete(*self.values_tree.get_children())
        for source in self._rows:
            if source.sheet != self.selected_sheet.get():
                continue
            item = f"row-{source.row}"
            target = self._alias_targets.get(source.location, "")
            if source.location in self._copied_locations:
                tag = ("aliased",)
            elif target:
                tag = ("skipped",)
                target = f"⚠ {target}"
            else:
                tag = () if source.eligible else ("unusable",)
            self.tree.insert(
                "", "end", iid=item,
                values=(target, source.row, source.name), tags=tag,
            )
            self.values_tree.insert(
                "", "end", iid=item,
                values=tuple(present(source.value_at(col)) for col in self.COLUMNS),
                tags=tag,
            )
            self._sources_by_item_id[item] = source
        self._notify_selection()

    def selected_source(self) -> SourceField | None:
        items = self.tree.selection() or self.values_tree.selection()
        return self._sources_by_item_id.get(items[0]) if items else None

    def source_at(self, tree: ttk.Treeview, y: int) -> SourceField | None:
        """Find the row under a mouse click, not the previously selected row."""
        if tree not in (self.tree, self.values_tree):
            return None
        return self._sources_by_item_id.get(tree.identify_row(y))

    def _on_sheet_changed(self, _event=None) -> None:
        self.refresh()

    def _notify_selection(self, _event=None) -> None:
        selected = self.selected_source()
        issue = self._alias_issues.get(selected.location) if selected else None
        self.issue_var.set(
            f"⚠ {selected.sheet}!B{selected.row}: {issue}" if issue else ""
        )
        if self._on_selection_changed is not None:
            self._on_selection_changed()
