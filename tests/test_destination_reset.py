"""Clean prior destination values without creating copies or touching layout."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from excel_copy_script.aliases import AliasStore
from excel_copy_script.gui import ExcelCopyApp
from excel_copy_script.workbook import WorkbookSession


class DestinationResetTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)
        self.path = self.directory / "summary.xlsx"
        self.aliases = AliasStore(self.directory / "aliases.json")

        book = Workbook()
        ws = book.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"], ws["D3"] = (
            "№", "Особовий склад", "Техніка", "Резерв",
        )
        ws["A4"], ws["A5"], ws["A6"] = 1, 2, 3
        ws["B4"] = 100
        ws["C5"] = "=10+20"
        ws["D6"] = "Історичні значення"
        ws["E5"] = "Поза стовпчиками імпорту"
        ws["B4"].fill = PatternFill(fill_type="solid", fgColor="00FFAA")
        extra = book.create_sheet("Додаткова інформація")
        extra["B4"] = "Не змінювати"
        book.save(self.path)
        self.session = WorkbookSession(self.path, self.aliases)

    def test_cleanup_saves_only_importable_columns_and_keeps_styles(self):
        self.assertEqual(self.session.clear_previous_imports(), 3)
        ws = load_workbook(self.path)["Дані_підрозділів"]
        self.assertIsNone(ws["B4"].value)
        self.assertIsNone(ws["C5"].value)
        self.assertIsNone(ws["D6"].value)
        self.assertEqual((ws["A4"].value, ws["A5"].value, ws["A6"].value), (1, 2, 3))
        self.assertEqual(ws["B3"].value, "Особовий склад")
        self.assertEqual(ws["C3"].value, "Техніка")
        self.assertEqual(ws["D3"].value, "Резерв")
        self.assertEqual(ws["E5"].value, "Поза стовпчиками імпорту")
        self.assertEqual(ws["B4"].fill.fgColor.rgb, "0000FFAA")
        self.assertEqual(
            load_workbook(self.path)["Додаткова інформація"]["B4"].value,
            "Не змінювати",
        )
        self.assertFalse(list(self.directory.glob("*.backup-*.xlsx")))
        # Empty import columns are available for a new run, regardless of A.
        self.assertEqual(self.session.next_empty_row(), 4)

    def test_second_cleanup_skips_save_and_does_not_create_backup(self):
        self.assertEqual(self.session.clear_previous_imports(), 3)
        with patch.object(self.session, "_save") as save:
            self.assertEqual(self.session.clear_previous_imports(), 0)
            save.assert_not_called()
        self.assertEqual(sorted(p.name for p in self.directory.iterdir()),
                         ["summary.xlsx"])

    def test_reset_after_previous_import_starts_again_at_row_four(self):
        self.session.clear_previous_imports()
        source = self.directory / "source.xlsx"
        book = Workbook()
        ws = book.active
        ws.title = "ВИСНОВКИ"
        ws["B1"], ws["C1"] = "Техніка", 41
        book.save(source)
        from excel_copy_script.models import SourceField
        self.aliases.set(
            self.session.template_id,
            SourceField(1, "Техніка", 41),
            self.session.fields[1],
        )
        self.assertEqual(self.session.open_input(source).row, 4)
        self.assertEqual(self.session.sheet["C4"].value, 41)
        second_session = WorkbookSession(self.path, AliasStore(self.aliases.path))
        self.assertEqual(second_session.clear_previous_imports(), 1)
        self.assertEqual(second_session.next_empty_row(), 4)
        self.assertEqual(second_session.open_input(source).row, 4)
        self.assertEqual(load_workbook(self.path)["Дані_підрозділів"]["C4"].value, 41)
        self.assertFalse(list(self.directory.glob("*.backup-*.xlsx")))

    def test_error_while_saving_restores_in_memory_and_disk_values(self):
        before = {"B4": self.session.sheet["B4"].value,
                  "C5": self.session.sheet["C5"].value,
                  "D6": self.session.sheet["D6"].value}
        with patch.object(self.session, "_save", side_effect=OSError("Disk full")):
            with self.assertRaisesRegex(OSError, "Disk full"):
                self.session.clear_previous_imports()
        for key, value in before.items():
            self.assertEqual(self.session.sheet[key].value, value)
        disk = load_workbook(self.path)["Дані_підрозділів"]
        for key, value in before.items():
            self.assertEqual(disk[key].value, value)
        self.assertFalse(list(self.directory.glob("*.backup-*.xlsx")))

    def test_gui_calls_cleanup_when_opening_destination(self):
        fake_session = SimpleNamespace(
            path=self.path,
            fields=[object()],
            clear_previous_imports=Mock(return_value=3),
            next_empty_row=Mock(return_value=4),
        )
        app = SimpleNamespace(
            recent_directories=SimpleNamespace(
                initial_directory=Mock(return_value=self.directory),
            ),
            session=None,
            aliases=None,
            _remember_directory=Mock(),
            destination_var=Mock(),
            source_var=Mock(),
            source_button=Mock(),
            clear_rows_var=Mock(get=Mock(return_value=True)),
            restart_button=Mock(),
            alias_button=Mock(),
            opened_sources=[Path("/before/old.xlsx")],
            _set_source_sheets=Mock(),
            _clear_tables=Mock(),
            _refresh_destination_rows=Mock(),
            status_var=Mock(),
        )
        with (
            patch("excel_copy_script.gui.filedialog.askopenfilename",
                  return_value=str(self.path)),
            patch("excel_copy_script.gui.AliasStore"),
            patch("excel_copy_script.gui.WorkbookSession", return_value=fake_session),
        ):
            ExcelCopyApp.open_destination(app)
        fake_session.clear_previous_imports.assert_called_once_with()
        self.assertEqual(app.opened_sources, [])
        app.restart_button.configure.assert_called_once_with(state="normal")
        self.assertIs(app.session, fake_session)
        app._remember_directory.assert_called_once_with("destination", str(self.path))

    def test_gui_does_not_switch_session_if_cleanup_fails(self):
        old_session = object()
        fake_session = SimpleNamespace(
            clear_previous_imports=Mock(side_effect=OSError("Disk full")),
        )
        app = SimpleNamespace(
            recent_directories=SimpleNamespace(
                initial_directory=Mock(return_value=self.directory),
            ),
            session=old_session,
            clear_rows_var=Mock(get=Mock(return_value=True)),
            _remember_directory=Mock(),
        )
        with (
            patch("excel_copy_script.gui.filedialog.askopenfilename",
                  return_value=str(self.path)),
            patch("excel_copy_script.gui.AliasStore"),
            patch("excel_copy_script.gui.WorkbookSession", return_value=fake_session),
            patch("excel_copy_script.gui.messagebox.showerror") as error_dialog,
        ):
            ExcelCopyApp.open_destination(app)
        error_dialog.assert_called_once()
        self.assertIs(app.session, old_session)
        app._remember_directory.assert_not_called()


if __name__ == "__main__":
    unittest.main()
