"""Real Tk regression checks for row-scoped alias creation and deletion."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.gui import ExcelCopyApp
from excel_copy_script.models import SourceField
from excel_copy_script.workbook import WorkbookSession


class RowAliasGuiTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        folder = Path(temp.name)
        destination = folder / "summary.xlsx"
        book = Workbook()
        ws = book.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"] = "№", "Резерв", "Техніка"
        book.save(destination)
        self.aliases = AliasStore(folder / "aliases.json")
        self.session = WorkbookSession(destination, self.aliases)
        self.source = folder / "report.xlsx"
        book = Workbook()
        ws = book.active
        ws.title = "ВИСНОВКИ"
        ws["B14"], ws["C14"], ws["F14"] = "Спільний", 18, 44
        book.save(self.source)
        self.session.open_input(self.source)

        self.app = ExcelCopyApp()
        self.app.geometry("1400x830")
        self.addCleanup(self.app.destroy)
        self.app.session = self.session
        self.app.aliases = self.aliases
        self.app.notebook.select(1)
        self.app._set_source_sheets(self.session.current.sheets)
        self.app._refresh_tables()
        self.app.update()

    def test_checked_row_only_creates_restricted_alias(self):
        self.app.create_sources.tree.selection_set("row-14")
        self.app.destination_tree.selection_set("3")
        self.app.row_only_var.set(True)
        self.app.create_sources.selected_value_column.set("F")
        self.app._on_selection()
        self.app.update()
        self.app.copy_button.invoke()
        self.assertEqual(
            self.aliases.get_exact(self.session.template_id, "Спільний", row=14)["source_column"],
            6,
        )
        self.assertIsNone(self.aliases.get_exact(self.session.template_id, "Спільний"))
        self.assertEqual(self.session.sheet["C4"].value, 44)
        self.assertEqual(self.session.current.copied[0].alias_row, 14)

    def test_unchecked_creates_general_alias(self):
        self.app.create_sources.tree.selection_set("row-14")
        self.app.destination_tree.selection_set("2")
        self.app._on_selection()
        self.app.update()
        self.app.copy_button.invoke()
        self.assertIsNotNone(self.aliases.get_exact(self.session.template_id, "Спільний"))
        self.assertIsNone(self.aliases.get_exact(self.session.template_id, "Спільний", row=14))

    def test_right_click_can_delete_only_general_when_scoped_also_exists(self):
        source = SourceField(14, "Спільний", 18)
        self.aliases.set(self.session.template_id, source, self.session.fields[0])
        self.aliases.set(self.session.template_id, source, self.session.fields[1], row=14)
        entries = self.app._aliases_for_source(source)
        self.assertEqual([e.source_row for e in entries], [None, 14])
        self.assertIsNotNone(self.app._alias_for_source(source, row=14))

        with patch("excel_copy_script.gui.messagebox.askyesno", return_value=True):
            self.app._delete_alias_entry(entries[0])
        self.assertIsNone(self.aliases.get_exact(self.session.template_id, "Спільний"))
        self.assertIsNotNone(
            self.aliases.get_exact(self.session.template_id, "Спільний", row=14),
        )

    def test_alias_manager_shows_scope_and_deletes_exact_variant(self):
        from excel_copy_script.alias_dialog import AliasManagerDialog

        source = SourceField(14, "Спільний", 18)
        self.aliases.set(self.session.template_id, source, self.session.fields[0])
        self.aliases.set(self.session.template_id, source, self.session.fields[1], row=14)
        # Avoid modal wait: instantiate dialog and exercise its list directly.
        dialog = AliasManagerDialog(
            self.app, self.aliases, self.session.template_id, self.session.fields,
        )
        self.addCleanup(lambda: dialog.winfo_exists() and dialog.destroy())
        self.app.update()
        rows = [dialog.tree.item(iid, "values") for iid in dialog.tree.get_children()]
        self.assertEqual(len(rows), 2)
        self.assertEqual({row[2] for row in rows}, {"Усі", "14"})
        scoped_id = next(iid for iid in dialog.tree.get_children()
                         if dialog.tree.item(iid, "values")[2] == "14")
        dialog.tree.selection_set(scoped_id)
        with patch("excel_copy_script.alias_dialog.messagebox.askyesno", return_value=True):
            dialog._delete_selected()
        self.assertIsNotNone(self.aliases.get_exact(self.session.template_id, "Спільний"))
        self.assertIsNone(
            self.aliases.get_exact(self.session.template_id, "Спільний", row=14),
        )
        dialog.destroy()


if __name__ == "__main__":
    unittest.main()
