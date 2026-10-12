"""Source columns C:G, whole-sheet preview and alias v3 regression tests."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.config import SOURCE_SHEET
from excel_copy_script.models import DestinationField, SourceField
from excel_copy_script.workbook import WorkbookSession


class SourceColumnTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.summary = self.folder / "summary.xlsx"
        self.alias_file = self.folder / "aliases.json"
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"] = "№", "Техніка", "Резерв"
        wb.save(self.summary)
        self.aliases = AliasStore(self.alias_file)
        self.session = WorkbookSession(self.summary, self.aliases)

    def make_input(self, filename, data):
        wb = Workbook()
        ws = wb.active
        ws.title = SOURCE_SHEET
        for location, value in data.items():
            ws[location] = value
        path = self.folder / filename
        wb.save(path)
        return path

    def test_manual_alias_uses_selected_f_column_and_reuses_it(self):
        first = self.make_input("first.xlsx", {
            "B2": "Бронетехніка", "C2": 1, "D2": 2, "E2": 3, "F2": 44, "G2": 5,
        })
        result = self.session.open_input(first)
        self.assertEqual(len(result.copied), 0)
        self.assertEqual(result.all_sources[1].values, (1, 2, 3, 44, 5))
        self.session.copy_manually(2, 2, source_column=6)
        self.assertEqual(self.session.sheet["B4"].value, 44)
        alias = self.aliases.get(self.session.template_id, "Бронетехніка")
        self.assertEqual(alias["source_column"], 6)

        second = self.make_input("second.xlsx", {
            "B3": "Бронетехніка", "C3": 999, "F3": 88,
        })
        result = self.session.open_input(second)
        self.assertEqual(len(result.copied), 1)
        self.assertEqual(result.copied[0].method, "Аліас")
        self.assertEqual(result.copied[0].source.value_column, 6)
        self.assertEqual(result.copied[0].source.value, 88)
        self.assertEqual(self.session.sheet["B5"].value, 88)

    def test_exact_header_names_do_not_copy_until_alias_is_created(self):
        path = self.make_input("exact.xlsx", {
            "B1": "Техніка", "D1": 90,
            "B2": "Резерв", "C2": 5, "G2": 6,
        })
        result = self.session.open_input(path)
        self.assertFalse(result.copied)
        self.assertEqual([x.row for x in result.unmatched_sources], [1, 2])
        self.assertIsNone(self.session.sheet["B4"].value)
        self.assertIsNone(self.session.sheet["C4"].value)

        # Despite both names matching output headers, the selected value
        # columns become usable only after manual alias creation.
        self.session.copy_manually(1, 2, source_column=4)
        self.session.copy_manually(2, 3, source_column=3)
        self.assertEqual(self.session.sheet["B4"].value, 90)
        self.assertEqual(self.session.sheet["C4"].value, 5)
        for name, column in (("Техніка", 4), ("Резерв", 3)):
            self.assertEqual(
                self.aliases.get(self.session.template_id, name)["source_column"],
                column,
            )

    def test_all_rows_available_even_without_valid_parameter(self):
        path = self.make_input("all.xlsx", {
            "B1": 42, "C1": 7,
            "B2": " ", "D2": 13,
            "B3": "Техніка", "C3": 100,
            "B4": "Без значення",
            "B5": "Можна вручну", "G5": 55,
        })
        self.aliases.set(
            self.session.template_id,
            SourceField(3, "Техніка", 100),
            self.session.fields[0],
        )
        result = self.session.open_input(path)
        self.assertEqual([row.row for row in result.all_sources], [1, 2, 3, 4, 5])
        self.assertEqual([row.eligible for row in result.all_sources],
                         [False, False, True, False, True])
        self.assertEqual([row.row for row in result.unmatched_sources], [5])
        self.assertTrue(any(r.source.row == 3 for r in result.copied))

    def test_bad_column_and_empty_cell_fail_without_data_change(self):
        path = self.make_input("missing.xlsx", {
            "B1": "Інша техніка", "G1": 17,
        })
        result = self.session.open_input(path)
        with self.assertRaises(ValueError):
            self.session.copy_manually(1, 2, source_column=5)
        with self.assertRaises(ValueError):
            self.session.copy_manually(1, 2, source_column=8)
        self.assertEqual(result.copied, [])
        self.assertIsNone(self.aliases.get(self.session.template_id, "Інша техніка"))
        self.session.copy_manually(1, 2, source_column=7)
        self.assertEqual(self.session.sheet["B4"].value, 17)

    def test_alias_for_other_sheet_reads_g_and_previews_destination(self):
        first = self.make_input("first.xlsx", {"B1": "Техніка", "C1": 10})
        self.aliases.set(
            self.session.template_id,
            SourceField(1, "Техніка", 10),
            self.session.fields[0],
        )
        wb = load_workbook(first)
        ws = wb.create_sheet("Склад")
        ws["B2"], ws["G2"] = "Нестача", 23
        wb.save(first)
        self.session.open_input(first)
        self.session.copy_manually(2, 3, sheet="Склад", source_column=7)
        self.assertEqual(self.session.source_alias_targets()[("склад", 2)], "C: Резерв")
        wb2 = load_workbook(first)
        wb2["Склад"]["G2"] = 77
        second = self.folder / "second.xlsx"
        wb2.save(second)
        result = self.session.open_input(second, row_selector=lambda request: None)
        self.assertEqual({r.method for r in result.copied}, {"Аліас"})
        self.assertEqual(self.session.sheet["C5"].value, 77)

    def test_v2_migration_sets_c_and_preserves_sheet_scope(self):
        old = {"version": 2, "templates": {
            self.session.template_id: {
                "висновки": {"старе": {"header": "техніка", "column": 2,
                                      "source_name": "Старе", "sheet_name": SOURCE_SHEET}},
                "склад": {"інше": {"header": "резерв", "column": 3,
                                   "source_name": "Інше", "sheet_name": "Склад"}},
            }
        }}
        original = json.dumps(old, ensure_ascii=False)
        self.alias_file.write_text(original, encoding="utf-8")
        updated = AliasStore(self.alias_file)
        self.assertEqual(updated.data["version"], 4)
        self.assertEqual(updated.get(self.session.template_id, "Старе")["source_column"], 3)
        self.assertEqual(updated.get(self.session.template_id, "Інше", "Склад")["source_column"], 3)
        self.assertEqual(updated.list_for_template(self.session.template_id)[0].source_column, 3)
        self.assertEqual((self.folder / "aliases.json.v2.bak").read_text(encoding="utf-8"), original)
        self.assertEqual(AliasStore(self.alias_file).data["version"], 4)

    def test_legacy_v1_migration_sets_c(self):
        old = {"version": 1, "templates": {self.session.template_id: {
            "старе": {"header": "техніка", "column": 2}
        }}}
        self.alias_file.write_text(json.dumps(old, ensure_ascii=False), encoding="utf-8")
        updated = AliasStore(self.alias_file)
        self.assertEqual(updated.get(self.session.template_id, "Старе")["source_column"], 3)
        self.assertIsNone(updated.get(self.session.template_id, "Старе", "Склад"))
        self.assertTrue((self.folder / "aliases.json.v1.bak").exists())


if __name__ == "__main__":
    unittest.main()
