"""Profile migration, isolation, changing Excel layouts, and last used selection."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.profiles import Profile, ProfileManager
from excel_copy_script.workbook import WorkbookSession
from excel_copy_script.workbook_layout import WorkbookLayout


class ProfileTests(unittest.TestCase):
    def setUp(self):
        d = tempfile.TemporaryDirectory()
        self.addCleanup(d.cleanup)
        self.home = Path(d.name)
        self.manager = ProfileManager(self.home)
        self.legacy = self.manager.legacy_path

    def test_first_start_migrates_aliases_without_modifying_original(self):
        old = {
            "version": 4,
            "templates": {"test-template": {
                "висновки": {"назва": {
                    "default": {
                        "header": "резерв", "column": 3, "source_name": "Назва",
                        "sheet_name": "ВИСНОВКИ", "source_column": 5,
                    },
                    "rows": {"14": {
                        "header": "резерв", "column": 3, "source_name": "Назва",
                        "sheet_name": "ВИСНОВКИ", "source_column": 7,
                    }},
                }},
            }},
        }
        self.home.mkdir(exist_ok=True)
        self.legacy.write_text(json.dumps(old, ensure_ascii=False), encoding="utf-8")
        original = self.legacy.read_bytes()
        profile = self.manager.load_last()
        self.assertEqual(profile.layout, WorkbookLayout())
        self.assertEqual(profile.aliases_data, old)
        self.assertEqual(self.legacy.read_bytes(), original)
        self.assertTrue(profile.path.exists())
        self.assertNotEqual(profile.path, self.legacy)
        self.assertEqual(Profile.load(profile.path).aliases_data, old)

    def test_v3_aliases_migrated_into_profile_without_touching_old_file(self):
        old = {"version": 3, "templates": {"abc": {"висновки": {
            "параметр": {
                "header": "ціль", "column": 2, "source_name": "Параметр",
                "sheet_name": "ВИСНОВКИ", "source_column": 6,
            }
        }}}}
        self.home.mkdir(exist_ok=True)
        self.legacy.write_text(json.dumps(old, ensure_ascii=False), encoding="utf-8")
        profile = self.manager.load_last()
        self.assertEqual(profile.aliases_data["version"], 4)
        self.assertEqual(profile.aliases_data["templates"]["abc"]["висновки"]["параметр"]["default"]["source_column"], 6)
        self.assertEqual(json.loads(self.legacy.read_text()), old)

    def test_invalid_legacy_migration_preserves_original_and_never_creates_empty_profile(self):
        self.home.mkdir(exist_ok=True)
        self.legacy.write_text("{corrupt", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Старий файл не змінено"):
            self.manager.load_last()
        self.assertEqual(self.legacy.read_text(), "{corrupt")
        self.assertFalse(self.manager.default_path.exists())

    def test_alias_modifications_save_inside_profile_not_legacy(self):
        self.home.mkdir(exist_ok=True)
        self.legacy.write_text('{"version": 4, "templates": {}}')
        profile = self.manager.load_last()
        from excel_copy_script.models import SourceField, DestinationField
        store = AliasStore.from_profile(profile)
        store.set("fingerprint", SourceField(1, "Значення", 5),
                  DestinationField(2, "Ціль"), row=1)
        loaded = Profile.load(profile.path)
        self.assertEqual(len(loaded.aliases_data["templates"]), 1)
        self.assertEqual(json.loads(self.legacy.read_text())["templates"], {})
        self.assertEqual(store.reload().get("fingerprint", "Значення", row=1)["column"], 2)

    def test_clone_remembers_last_profile_and_preserves_independent_aliases(self):
        first = self.manager.load_last()
        other = first.copy_as(self.home / "other-profile.json")
        self.manager.remember(other)
        self.assertEqual(self.manager.load_last().path, other.path)
        from excel_copy_script.models import DestinationField, SourceField
        AliasStore.from_profile(other).set(
            "template", SourceField(3, "Інший", 11),
            DestinationField(2, "Результат"),
        )
        self.assertNotIn("template", Profile.load(first.path).aliases_data["templates"])
        with self.assertRaises(FileExistsError):
            first.copy_as(other.path)

    def test_unavailable_last_profile_fails_without_overwriting_state(self):
        first = self.manager.load_last()
        other = first.copy_as(self.home / "other.json")
        self.manager.remember(other)
        other.path.unlink()
        with self.assertRaisesRegex(ValueError, "останній профіль"):
            self.manager.load_last()
        self.assertTrue(first.path.exists())
        self.assertIn("other.json", self.manager.state_path.read_text())

    def test_nondefault_excel_columns_sheet_and_input_rows_are_active(self):
        layout = WorkbookLayout(
            destination_sheet="Custom",
            header_row=6, first_data_row=8, first_destination_column=3,
            source_name_column=4, first_source_row=5,
            source_value_first_column=6, source_value_last_column=8,
            preferred_source_sheet="Input",
        )
        profile = self.manager.load_last()
        profile.layout = layout
        profile.save()
        self.assertEqual(Profile.load(profile.path).layout, layout)

        output_path = self.home / "out.xlsx"
        output = Workbook()
        output.active.title = "Custom"
        ws = output.active
        ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "Ordinal", "Ignored", "Amount", "More"
        ws["A8"], ws["B8"] = 1, "keep"
        output.save(output_path)
        store = AliasStore.from_profile(Profile.load(profile.path))
        session = WorkbookSession(output_path, store)
        self.assertEqual([(field.column, field.name) for field in session.fields],
                         [(3, "Amount"), (4, "More")])
        self.assertEqual(session.next_empty_row(), 8)
        source_path = self.home / "input.xlsx"
        source = Workbook()
        source.active.title = "Input"
        ws = source.active
        ws["D2"], ws["F2"] = "Ігнорувати", 1000
        ws["D5"], ws["F5"], ws["G5"], ws["H5"] = "Рахунок", 10, 20, 30
        source.save(source_path)
        from excel_copy_script.models import SourceField
        store.set(session.template_id, SourceField(
            5, "Рахунок", 30, sheet="Input", value_first_column=6,
            value_column=8,
        ), session.fields[0])
        result = session.open_input(source_path)
        self.assertEqual(result.row, 8)
        self.assertEqual(result.copied[0].source.value, 30)
        self.assertEqual(session.sheet["C8"].value, 30)
        self.assertEqual(session.sheet["B8"].value, "keep")
        self.assertEqual(session.sheet["C7"].value, None)
        self.assertFalse(any(f.row == 2 for f in result.all_sources))
        session.clear_previous_imports()
        self.assertIsNone(session.sheet["C8"].value)
        self.assertEqual(session.sheet["B8"].value, "keep")

    def test_create_empty_profile_is_independent_from_existing_aliases(self):
        first = self.manager.load_last()
        from excel_copy_script.models import DestinationField, SourceField
        AliasStore.from_profile(first).set(
            "original", SourceField(4, "Old", 1), DestinationField(2, "Data"),
        )
        destination = self.home / "new-blank.json"
        layout = WorkbookLayout(
            destination_sheet="Custom", header_row=5, first_data_row=7,
            first_destination_column=4, source_name_column=2,
            first_source_row=6, source_value_first_column=3,
            source_value_last_column=5, preferred_source_sheet="Source",
        )
        new_profile = Profile.create_empty(destination, layout)
        self.assertEqual(new_profile.aliases_data["templates"], {})
        self.assertEqual(new_profile.layout, layout)
        self.assertEqual(Profile.load(destination).aliases_data["templates"], {})
        self.assertIn(
            "original", Profile.load(first.path).aliases_data["templates"],
        )
        self.manager.remember(new_profile)
        self.assertEqual(self.manager.load_last().path, destination.resolve())

    def test_create_empty_profile_refuses_existing_file_without_changing_it(self):
        existing = self.home / "existing.json"
        existing.write_text('{"custom": true}', encoding="utf-8")
        before = existing.read_bytes()
        with self.assertRaises(FileExistsError):
            Profile.create_empty(existing)
        self.assertEqual(existing.read_bytes(), before)

    def test_create_empty_profile_validates_layout_before_writing(self):
        target = self.home / "bad.json"
        with self.assertRaisesRegex(ValueError, "Перший рядок даних"):
            Profile.create_empty(
                target, WorkbookLayout(header_row=8, first_data_row=7),
            )
        self.assertFalse(target.exists())

    def test_malformed_layout_rejected(self):
        p = self.manager.load_last()
        data = json.loads(p.path.read_text())
        data["layout"]["source_name_column"] = 6
        p.path.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Колонка назви"):
            Profile.load(p.path)


if __name__ == "__main__":
    unittest.main()
