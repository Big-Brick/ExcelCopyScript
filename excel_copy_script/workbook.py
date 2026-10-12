"""ExcelCopyScript: interactive transfer from ВИСНОВКИ to Дані_підрозділів."""

from __future__ import annotations

import os
import shutil
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .aliases import AliasStore
from .config import DESTINATION_SHEET, FIRST_DATA_ROW, FIRST_DESTINATION_COLUMN, HEADER_ROW, SOURCE_SHEET
from .models import CopyRecord, DestinationField, ImportResult, SourceField
from .utils import nonempty, text_as_excel_value

class WorkbookSession:
    def __init__(self, path: Path, aliases: AliasStore):
        self.path = Path(path)
        if self.path.suffix.lower() != ".xlsx":
            raise ValueError("Підтримуються лише файли .xlsx")
        self.workbook = load_workbook(self.path)
        if DESTINATION_SHEET not in self.workbook:
            raise ValueError(f"У підсумковому файлі немає аркуша «{DESTINATION_SHEET}»")
        self.sheet = self.workbook[DESTINATION_SHEET]
        header_fields = [DestinationField(cell.column, cell.value.strip())
                         for cell in self.sheet[HEADER_ROW]
                         if isinstance(cell.value, str) and cell.value.strip()]
        # Column A contains a row number, not a field to import or match.
        self.fields = [field for field in header_fields
                       if field.column >= FIRST_DESTINATION_COLUMN]
        if not self.fields:
            raise ValueError(
                f"У рядку {HEADER_ROW}, починаючи зі стовпчика B, немає текстових заголовків"
            )
        self.aliases = aliases
        self.template_id = aliases.template_id(self.fields)
        # Older releases included column A in the alias template fingerprint.
        # Migrate their saved mappings without ever restoring mappings to A.
        legacy_template_id = aliases.template_id(header_fields)
        if legacy_template_id != self.template_id:
            aliases.migrate_legacy_template(
                legacy_template_id, self.template_id, self.fields
            )
        self.current: ImportResult | None = None
        self.backup_path: Path | None = None

    def next_empty_row(self) -> int:
        # An ordinal number in A does not make a row occupied. Still inspect
        # every other column to avoid overwriting partially populated records.
        for row in range(FIRST_DATA_ROW, self.sheet.max_row + 1):
            if all(not nonempty(self.sheet.cell(row, col).value)
                   for col in range(FIRST_DESTINATION_COLUMN, self.sheet.max_column + 1)):
                return row
        return max(FIRST_DATA_ROW, self.sheet.max_row + 1)

    def _read_source(self, path: Path) -> tuple[list[SourceField], list[str], list[str]]:
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Підтримуються лише файли .xlsx")
        if path.resolve() == self.path.resolve():
            raise ValueError("Вхідний і підсумковий файли не можуть збігатися")
        cached = load_workbook(path, read_only=True, data_only=True)
        formulas = None
        try:
            if SOURCE_SHEET not in cached:
                raise ValueError(f"У файлі {path.name} немає аркуша «{SOURCE_SHEET}»")
            # Read formula metadata separately: openpyxl does not calculate formulas.
            formulas = load_workbook(path, read_only=True, data_only=False)
            entries: list[SourceField] = []
            warnings: list[str] = []
            # Process ВИСНОВКИ first, then all other worksheets. All worksheets
            # use the same B (field name) / C (cached value) layout.
            sheet_names = [SOURCE_SHEET] + [name for name in cached.sheetnames
                                           if name != SOURCE_SHEET]
            for sheet_name in sheet_names:
                cached_ws = cached[sheet_name]
                formula_ws = formulas[sheet_name]
                for row, (name, cached_value) in enumerate(
                    cached_ws.iter_rows(min_col=2, max_col=3, values_only=True), start=1
                ):
                    if not isinstance(name, str) or not name.strip():
                        continue
                    if not nonempty(cached_value):
                        formula_cell = formula_ws.cell(row, 3)
                        if formula_cell.data_type == "f":
                            warnings.append(
                                f"{sheet_name}!C{row}: формула без збереженого результату — пропущено"
                            )
                        continue
                    entries.append(SourceField(row, name.strip(), cached_value, sheet=sheet_name))
            return entries, warnings, sheet_names
        finally:
            cached.close()
            if formulas is not None:
                formulas.close()

    def _destination_for(self, source: SourceField) -> tuple[DestinationField | None, str]:
        alias = self.aliases.get(self.template_id, source.name, source.sheet)
        if alias:
            candidates = [field for field in self.fields if field.key == alias.get("header")]
            if len(candidates) == 1:
                return candidates[0], "Аліас"
            if len(candidates) > 1:
                pinned = [field for field in candidates if field.column == alias.get("column")]
                if len(pinned) == 1:
                    return pinned[0], "Аліас"
        if source.sheet == SOURCE_SHEET:
            exact = [field for field in self.fields if field.key == source.key]
            if len(exact) == 1:
                return exact[0], "Збіг назви"
        return None, ""

    def _write(self, row: int, records: list[CopyRecord]) -> None:
        if not records:
            return
        before = {}
        for record in records:
            column = record.destination.column
            if column < FIRST_DESTINATION_COLUMN:
                raise ValueError("Стовпчик A містить порядкові номери і не призначений для імпорту")
            cell = self.sheet.cell(row, column)
            if nonempty(cell.value):
                raise ValueError(f"Комірка {cell.coordinate} вже заповнена")
            before[column] = cell.value
            cell.value = text_as_excel_value(record.source.value)
        try:
            self._save()
        except Exception:
            for column, old_value in before.items():
                self.sheet.cell(row, column).value = old_value
            raise

    def _save(self) -> None:
        if self.backup_path is None:
            suffix = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
            backup = self.path.with_name(f"{self.path.stem}.backup-{suffix}.xlsx")
            shutil.copy2(self.path, backup)
            self.backup_path = backup
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(dir=self.path.parent, prefix=".excel-copy-",
                                             suffix=".xlsx", delete=False) as file:
                tmp_path = Path(file.name)
            self.workbook.save(tmp_path)
            os.chmod(tmp_path, self.path.stat().st_mode)
            os.replace(tmp_path, self.path)
        finally:
            if tmp_path is not None:
                tmp_path.unlink(missing_ok=True)

    def open_input(self, path: Path) -> ImportResult:
        source_fields, warnings, sheet_names = self._read_source(Path(path))
        target_row = self.next_empty_row()
        duplicates = Counter((field.sheet.casefold(), field.key) for field in source_fields)
        proposals: dict[int, list[CopyRecord]] = defaultdict(list)
        for source in source_fields:
            if duplicates[(source.sheet.casefold(), source.key)] > 1:
                continue
            destination, method = self._destination_for(source)
            if destination is not None:
                proposals[destination.column].append(CopyRecord(source, destination, method))
        safe_matches = [matches[0] for matches in proposals.values() if len(matches) == 1]
        safe_matches.sort(key=lambda rec: rec.destination.column)
        self._write(target_row, safe_matches)
        result = ImportResult(filename=Path(path).name, row=target_row, copied=safe_matches,
                              warnings=warnings, sheets=sheet_names)
        if any(count > 1 for count in duplicates.values()):
            result.warnings.append("Повторювані назви у вхідному файлі не зіставлено автоматично")
        if any(len(matches) > 1 for matches in proposals.values()):
            result.warnings.append("Кілька вхідних параметрів претендують на один стовпчик — пропущено")
        used_rows = {record.source.location for record in safe_matches}
        used_columns = {record.destination.column for record in safe_matches}
        result.unmatched_sources = [source for source in source_fields
                                    if source.location not in used_rows]
        result.unmatched_destinations = [dest for dest in self.fields if dest.column not in used_columns]
        self.current = result
        return result

    def copy_manually(
        self, source_row: int, destination_col: int, sheet: str = SOURCE_SHEET
    ) -> str | None:
        result = self.current
        if result is None:
            raise ValueError("Спочатку відкрийте вхідний файл")
        source = next((s for s in result.unmatched_sources
                       if s.row == source_row and s.sheet.casefold() == sheet.casefold()), None)
        destination = next((d for d in result.unmatched_destinations if d.column == destination_col), None)
        if source is None or destination is None:
            raise ValueError("Вибраний параметр уже був скопійований")
        record = CopyRecord(source, destination, "Вручну")
        self._write(result.row, [record])
        result.copied.append(record)
        result.unmatched_sources.remove(source)
        result.unmatched_destinations.remove(destination)
        try:
            self.aliases.set(self.template_id, source, destination)
        except (OSError, ValueError) as exc:
            return f"Значення скопійовано, але аліас не збережено: {exc}"
        return None
