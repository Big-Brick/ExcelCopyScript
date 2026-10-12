"""Tkinter application shell: reviewing imported values and creating aliases."""

from __future__ import annotations

import tkinter as tk
from pathlib import Path

from openpyxl.utils import get_column_letter
from tkinter import filedialog, messagebox, ttk

from .alias_dialog import AliasManagerDialog
from .aliases import AliasStore
from .config import FILE_TYPES
from .gui_layout import LayoutPersistence
from .models import AliasEntry, CopyRecord, SourceField
from .recent_directories import RecentDirectories
from .source_rows_panel import SourceRowsPanel
from .utils import nonempty, present
from .workbook import WorkbookSession


class ExcelCopyApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Перенесення даних Excel")
        self.geometry("1260x720")
        self.minsize(900, 510)
        self.session: WorkbookSession | None = None
        self.aliases: AliasStore | None = None
        self.recent_directories = RecentDirectories()
        self.destination_var = tk.StringVar(value="Підсумковий файл не вибрано")
        self.source_var = tk.StringVar(value="Вхідний файл не вибрано")
        self.status_var = tk.StringVar(value="Виберіть підсумковий файл, потім вхідний.")
        self._layout_error_reported = False
        self._build_ui()
        self.layout = LayoutPersistence(on_error=self._on_layout_error)
        self._configure_layout_persistence()
        self.notebook.bind("<<NotebookTabChanged>>", self._on_notebook_tab_changed)
        self.after_idle(self._restore_active_tab_layout)

    def _build_ui(self) -> None:
        top = ttk.Frame(self, padding=12)
        top.pack(fill="x")
        ttk.Button(
            top, text="1. Відкрити підсумковий файл", command=self.open_destination,
        ).pack(side="left")
        self.source_button = ttk.Button(
            top, text="2. Відкрити вхідний / наступний файл",
            command=self.open_source, state="disabled",
        )
        self.source_button.pack(side="left", padx=(12, 0))
        ttk.Label(self, textvariable=self.destination_var, padding=(12, 0)).pack(anchor="w")
        ttk.Label(self, textvariable=self.source_var, padding=(12, 2)).pack(anchor="w")

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True, padx=12, pady=(8, 4))
        review_tab = ttk.Frame(self.notebook)
        create_tab = ttk.Frame(self.notebook)
        self.notebook.add(review_tab, text="Аналіз і видалення аліасів")
        self.notebook.add(create_tab, text="Створення аліасів")

        # Review: copied rows + unmatched source rows (independent sheet selector).
        review_panes = ttk.Panedwindow(review_tab, orient="horizontal")
        self.review_panes = review_panes
        review_panes.pack(fill="both", expand=True)
        copied_frame, self.copied_tree = self._make_table(
            review_panes, "Скопійовані значення",
            ("Звідки", "Куди", "Значення", "Спосіб"), (180, 190, 140, 130),
        )
        self.review_sources = SourceRowsPanel(review_panes)
        review_panes.add(copied_frame, weight=1)
        review_panes.add(self.review_sources, weight=1)

        review_actions = ttk.Frame(review_tab, padding=(4, 8, 4, 6))
        review_actions.pack(fill="x")
        self.delete_alias_button = ttk.Button(
            review_actions, text="Видалити аліас вибраного запису",
            command=self.delete_selected_alias, state="disabled",
        )
        self.delete_alias_button.pack(side="left")
        self.alias_button = ttk.Button(
            review_actions, text="Керувати всіма аліасами",
            command=self.open_alias_manager, state="disabled",
        )
        self.alias_button.pack(side="left", padx=(12, 0))
        ttk.Label(
            review_tab,
            text="Видалення аліаса не змінює вже записані значення Excel. "
                 "Для створення нового аліаса перейдіть на другу вкладку.",
            padding=(4, 0, 4, 8),
        ).pack(anchor="w")
        self.copied_tree.bind("<<TreeviewSelect>>", self._on_copied_selection)

        # Creation: all destination fields + all source rows; completed
        # destinations are visible for context but cannot be selected for copying.
        create_panes = ttk.Panedwindow(create_tab, orient="horizontal")
        self.create_panes = create_panes
        create_panes.pack(fill="both", expand=True)
        destination_frame, self.destination_tree = self._make_table(
            create_panes, "Вихідні стовпчики",
            ("Колонка", "Назва", "Джерело значення"), (80, 220, 320),
        )
        self.destination_tree.tag_configure("filled", background="#ddf2df")
        self.create_sources = SourceRowsPanel(
            create_panes, on_selection_changed=self._on_selection,
        )
        create_panes.add(destination_frame, weight=1)
        create_panes.add(self.create_sources, weight=1)
        self.destination_tree.bind("<<TreeviewSelect>>", self._on_selection)
        self.destination_tree.bind("<Button-3>", self._show_destination_alias_menu)
        for tree in (self.create_sources.tree, self.create_sources.values_tree):
            tree.bind("<Button-3>", self._show_source_alias_menu)

        create_actions = ttk.Frame(create_tab, padding=(4, 8, 4, 6))
        create_actions.pack(fill="x")
        self.copy_button = ttk.Button(
            create_actions, text="Копіювати вибрану пару + запам'ятати аліас",
            command=self.copy_selected, state="disabled",
        )
        self.copy_button.pack(anchor="w")
        ttk.Label(
            create_tab,
            text="Виберіть вихідний стовпчик, вхідний рядок і колонку значення C–G.",
            padding=(4, 0, 4, 8),
        ).pack(anchor="w")

        status = ttk.Frame(self, padding=(12, 4, 12, 12))
        status.pack(fill="x")
        ttk.Separator(status).pack(fill="x", pady=(0, 8))
        ttk.Label(status, textvariable=self.status_var, wraplength=1150).pack(anchor="w")

    def _on_layout_error(self, error: OSError) -> None:
        if self._layout_error_reported:
            return
        self._layout_error_reported = True
        messagebox.showwarning(
            "Налаштування інтерфейсу",
            f"Не вдалося зберегти ширини колонок і розділювачі:\n{error}",
            parent=self,
        )

    def _configure_layout_persistence(self) -> None:
        trees = {
            "review_copied": self.copied_tree,
            "review_source_headers": self.review_sources.tree,
            "review_source_values": self.review_sources.values_tree,
            "create_destination": self.destination_tree,
            "create_source_headers": self.create_sources.tree,
            "create_source_values": self.create_sources.values_tree,
        }
        for key, tree in trees.items():
            self.layout.register_columns(key, tree)
        for tab, key, splitter in (
            ("review", "review_outer", self.review_panes),
            ("review", "review_inner", self.review_sources.divider),
            ("create", "create_outer", self.create_panes),
            ("create", "create_inner", self.create_sources.divider),
        ):
            self.layout.register_sash(tab, key, splitter)

    def _restore_active_tab_layout(self, _event=None) -> None:
        tab = "review" if self.notebook.index("current") == 0 else "create"
        self.layout.restore_tab(tab)

    def _on_notebook_tab_changed(self, _event=None) -> None:
        # The page must be mapped before positioning ttk.Panedwindow sashes.
        self.after_idle(self._restore_active_tab_layout)

    @staticmethod
    def _make_table(parent, title, headings, widths):
        frame = ttk.Frame(parent, padding=4)
        ttk.Label(frame, text=title, font=("", 10, "bold")).pack(anchor="w", pady=(0, 5))
        container = ttk.Frame(frame)
        container.pack(fill="both", expand=True)
        tree = ttk.Treeview(container, columns=headings, show="headings", selectmode="browse")
        for heading, width in zip(headings, widths):
            tree.heading(heading, text=heading)
            tree.column(heading, width=width, minwidth=60, stretch=True)
        scrollbar_y = ttk.Scrollbar(container, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar_y.set)
        tree.pack(side="left", fill="both", expand=True)
        scrollbar_y.pack(side="right", fill="y")
        return frame, tree

    def _on_selection(self, _event=None) -> None:
        can_copy = bool(
            self.session and self.session.current
            and self.destination_tree.selection()
            and int(self.destination_tree.selection()[0]) in {
                field.column for field in self.session.current.unmatched_destinations
            }
            and (source := self.create_sources.selected_source()) is not None
            and source.eligible
            and source.location in {
                row.location for row in self.session.current.unmatched_sources
            }
            and nonempty(source.value_at(self.create_sources.source_column()))
        )
        self.copy_button.configure(state="normal" if can_copy else "disabled")

    def _selected_copy_record(self) -> CopyRecord | None:
        selected = self.copied_tree.selection()
        if not selected or self.session is None or self.session.current is None:
            return None
        try:
            return self.session.current.copied[int(selected[0])]
        except (IndexError, ValueError):
            return None

    def _record_has_alias(self, record: CopyRecord) -> bool:
        """Only allow deleting an alias that resolves to this record's destination."""
        if self.aliases is None or self.session is None:
            return False
        alias = self.aliases.get(
            self.session.template_id, record.source.name, record.source.sheet,
        )
        if (not alias or alias.get("header") != record.destination.key
                or alias.get("source_column", 3) != record.source.value_column):
            return False
        duplicates = sum(field.key == record.destination.key for field in self.session.fields)
        return duplicates == 1 or alias.get("column") == record.destination.column

    def _on_copied_selection(self, _event=None) -> None:
        record = self._selected_copy_record()
        enabled = record is not None and self._record_has_alias(record)
        self.delete_alias_button.configure(state="normal" if enabled else "disabled")

    def _alias_for_source(self, source: SourceField) -> AliasEntry | None:
        if self.aliases is None or self.session is None:
            return None
        return next((
            entry for entry in self.aliases.list_for_template(self.session.template_id)
            if entry.source_key == source.key
            and entry.sheet.casefold() == source.sheet.casefold()
        ), None)

    def _delete_alias_entry(self, entry: AliasEntry) -> None:
        """Delete one saved mapping without altering values already copied."""
        if self.aliases is None or self.session is None:
            return
        if not messagebox.askyesno(
            "Підтвердження видалення",
            f"Видалити аліас «{entry.sheet}!{entry.source_name}» "
            f"[{get_column_letter(entry.source_column)}] → "
            f"«{entry.destination_key}»?\n\n"
            "Уже скопійовані значення Excel не зміняться.",
            parent=self,
        ):
            return
        try:
            deleted = self.aliases.delete(
                self.session.template_id, entry.source_key, sheet=entry.sheet,
            )
        except (OSError, ValueError) as exc:
            messagebox.showerror("Помилка видалення аліаса", str(exc), parent=self)
            return
        if deleted:
            self._refresh_copied()
            self._refresh_source_panels()

    def _popup_alias_menu(self, entries: list[AliasEntry], event) -> None:
        if not entries:
            return
        # Retain a live Tk menu after tk_popup() returns (selection is async).
        previous = getattr(self, "_alias_context_menu", None)
        if previous is not None:
            previous.destroy()
        menu = tk.Menu(self, tearoff=False)
        self._alias_context_menu = menu
        if len(entries) == 1:
            entry = entries[0]
            menu.add_command(
                label="Видалити аліас",
                command=lambda entry=entry: self._delete_alias_entry(entry),
            )
        else:
            submenu = tk.Menu(menu, tearoff=False)
            for entry in entries:
                menu_label = (
                    f"{entry.sheet}!{entry.source_name} "
                    f"[{get_column_letter(entry.source_column)}]"
                )
                submenu.add_command(
                    label=menu_label,
                    command=lambda entry=entry: self._delete_alias_entry(entry),
                )
            menu.add_cascade(label="Видалити аліас", menu=submenu)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _show_source_alias_menu(self, event):
        if self.aliases is None or self.session is None:
            return
        source = self.create_sources.source_at(event.widget, event.y)
        if source is None:
            return
        entry = self._alias_for_source(source)
        if entry is not None:
            self._popup_alias_menu([entry], event)
            return "break"

    def _show_destination_alias_menu(self, event):
        if self.session is None:
            return
        row_id = self.destination_tree.identify_row(event.y)
        if not row_id:
            return
        entries = self.session.aliases_for_destination(int(row_id))
        if entries:
            self._popup_alias_menu(entries, event)
            return "break"

    def delete_selected_alias(self) -> None:
        """Legacy review-tab action, sharing the same confirmation and refresh."""
        record = self._selected_copy_record()
        if record is None or not self._record_has_alias(record):
            return
        entry = self._alias_for_source(record.source)
        if entry is not None:
            self._delete_alias_entry(entry)

    def _remember_directory(self, kind: str, filename: str) -> None:
        """Do not fail a successful file open because UI preferences cannot be saved."""
        try:
            self.recent_directories.remember_file(kind, filename)
        except (OSError, ValueError) as exc:
            messagebox.showwarning(
                "Не вдалося запам'ятати теку",
                f"Файл відкрито, але теку не вдалося зберегти:\n{exc}",
            )

    def open_destination(self) -> None:
        chosen = filedialog.askopenfilename(
            title="Підсумковий Excel-файл", filetypes=FILE_TYPES,
            initialdir=str(self.recent_directories.initial_directory("destination")),
        )
        if not chosen:
            return
        try:
            aliases = AliasStore()
            session = WorkbookSession(Path(chosen), aliases)
            session.clear_previous_imports()
        except Exception as exc:
            messagebox.showerror("Помилка відкриття", str(exc))
            return
        self.aliases, self.session = aliases, session
        self._remember_directory("destination", chosen)
        self.destination_var.set(f"Підсумковий: {session.path}")
        self.source_var.set("Вхідний файл не вибрано")
        self.source_button.configure(state="normal")
        self.alias_button.configure(state="normal")
        self._set_source_sheets([])
        self._clear_tables()
        self._refresh_destination_rows()
        self.status_var.set(
            f"Знайдено {len(session.fields)} стовпчиків. "
            f"Наступний рядок: {session.next_empty_row()}."
        )

    def open_alias_manager(self) -> None:
        if self.session is None or self.aliases is None:
            return
        dialog = AliasManagerDialog(
            self, self.aliases, self.session.template_id, self.session.fields,
            layout=self.layout,
        )
        self.wait_window(dialog)
        self._refresh_copied()
        self._refresh_source_panels()

    def open_source(self) -> None:
        if self.session is None:
            return
        chosen = filedialog.askopenfilename(
            title="Вхідний Excel-файл", filetypes=FILE_TYPES,
            initialdir=str(self.recent_directories.initial_directory(
                "source", fallback=self.session.path.parent,
            )),
        )
        if not chosen:
            return
        try:
            result = self.session.open_input(Path(chosen))
        except Exception as exc:
            messagebox.showerror("Помилка імпорту", str(exc))
            return
        self._remember_directory("source", chosen)
        self.source_var.set(f"Вхідний: {chosen} → рядок {result.row}")
        self._set_source_sheets(result.sheets)
        self._refresh_tables()

    def _set_source_sheets(self, sheets: list[str]) -> None:
        for panel in (self.review_sources, self.create_sources):
            panel.set_sheets(sheets)

    def _clear_tables(self) -> None:
        for tree in (self.copied_tree, self.destination_tree):
            tree.delete(*tree.get_children())
        for panel in (self.review_sources, self.create_sources):
            panel.refresh([], alias_targets={})
        self._on_selection()
        self._on_copied_selection()

    def _refresh_copied(self) -> None:
        self.copied_tree.delete(*self.copied_tree.get_children())
        result = self.session.current if self.session is not None else None
        if result is not None:
            for index, record in enumerate(result.copied):
                method = record.method
                if not self._record_has_alias(record):
                    method = "Аліас видалено"
                self.copied_tree.insert(
                    "", "end", iid=str(index),
                    values=(
                        (f"{record.source.sheet}!B{record.source.row}: {record.source.name} "
                         f"[{get_column_letter(record.source.value_column)}{record.source.row}]"),
                        f"{record.destination.letter}: {record.destination.name}",
                        present(record.source.value), method,
                    ),
                )
        self._on_copied_selection()

    def _refresh_destination_rows(self) -> None:
        """Show every output column, marking values copied in the current import."""
        self.destination_tree.delete(*self.destination_tree.get_children())
        if self.session is None:
            self._on_selection()
            return

        result = self.session.current
        copies = {
            record.destination.column: record
            for record in result.copied
        } if result is not None else {}
        for destination in self.session.fields:
            record = copies.get(destination.column)
            origin = ""
            if record is not None:
                source = record.source
                origin = (
                    f"{source.sheet}!{get_column_letter(source.value_column)}"
                    f"{source.row}: {source.name}"
                )
            self.destination_tree.insert(
                "", "end", iid=str(destination.column),
                values=(destination.letter, destination.name, origin),
                tags=("filled",) if record is not None else (),
            )
        self._on_selection()

    def _refresh_source_panels(self) -> None:
        result = self.session.current if self.session is not None else None
        rows = result.all_sources if result is not None else []
        targets = self.session.source_alias_targets() if result is not None else {}
        for panel in (self.review_sources, self.create_sources):
            panel.refresh(rows, alias_targets=targets)

    def _refresh_tables(self) -> None:
        self._clear_tables()
        if self.session is None or self.session.current is None:
            return
        result = self.session.current
        self._refresh_copied()
        self._refresh_destination_rows()
        self._refresh_source_panels()
        self._on_selection()
        status = (
            f"{result.filename} → рядок {result.row}. Скопійовано: {len(result.copied)}; "
            f"не зіставлено вихідних стовпчиків: {len(result.unmatched_destinations)}; "
            f"вхідних рядків: {len(result.unmatched_sources)}. Зміни автоматично збережено."
        )
        if result.warnings:
            status += " Увага: " + "; ".join(result.warnings[:4])
        self.status_var.set(status)

    def copy_selected(self) -> None:
        if self.session is None:
            return
        source = self.create_sources.selected_source()
        destinations = self.destination_tree.selection()
        if source is None or not destinations:
            return
        if int(destinations[0]) not in {
            field.column for field in self.session.current.unmatched_destinations
        }:
            return  # Already copied; visible for context, but not writable again.
        if not source.eligible or source.location not in {
            row.location for row in self.session.current.unmatched_sources
        }:
            messagebox.showwarning(
                "Неможливо скопіювати", "Виберіть рядок із назвою та доступним значенням.",
            )
            return
        try:
            warning = self.session.copy_manually(
                source.row, int(destinations[0]), sheet=source.sheet,
                source_column=self.create_sources.source_column(),
            )
        except Exception as exc:
            messagebox.showerror("Помилка копіювання", str(exc))
            return
        self._refresh_tables()
        if warning:
            messagebox.showwarning("Аліас не збережено", warning)


def main() -> None:
    ExcelCopyApp().mainloop()
