"""Regression tests for imported percentages and aliased rows skipped silently."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.models import SourceField
from excel_copy_script.workbook import WorkbookSession


PARAMETER = "Загальна результативність виконання завдань"


class ImportDiagnosticsTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.folder = Path(tmp.name)
        destination = self.folder / "destination.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        for col, name in ("A", "№"), ("B", "Склад"), ("C", "Техніка"), ("D", "Резерв"), ("E", "Результативність"):
            ws[f"{col}3"] = name
        wb.save(destination)
        self.aliases = AliasStore(self.folder / "aliases.json")
        self.session = WorkbookSession(destination, self.aliases)
        self.destination = self.session.fields[-1]

    def make_source(self, second=None, text_percent=False):
        wb = Workbook()
        ws = wb.active
        ws.title = "ВИСНОВКИ"
        ws["A14"] = 14
        ws["B14"] = PARAMETER
        ws["C14"] = "89%" if text_percent else 0.89
        if not text_percent:
            ws["C14"].number_format = "0%"
        ws["D14"] = "80%"
        ws["E14"] = "П1"
        ws["F14"] = "у межах норми"
        if second is not None:
            row, name, c_value, d_value = second
            ws[f"B{row}"] = name
            if c_value is not None:
                ws[f"C{row}"] = c_value
            if d_value is not None:
                ws[f"D{row}"] = d_value
        path = self.folder / "input.xlsx"
        wb.save(path)
        return path

    def set_alias(self, name=PARAMETER, column=3):
        self.aliases.set(
            self.session.template_id,
            SourceField(14, name, 0.89, value_column=column),
            self.destination,
        )

    def test_numeric_excel_percentage_is_copied_to_destination_e(self):
        self.set_alias()
        result = self.session.open_input(self.make_source())
        self.assertEqual(len(result.copied), 1)
        self.assertEqual(result.copied[0].source.row, 14)
        self.assertEqual(result.copied[0].destination.letter, "E")
        self.assertEqual(result.copied[0].source.value, 0.89)
        self.assertFalse(result.alias_issues)
        self.assertEqual(
            self.session.source_alias_targets()[("висновки", 14)],
            "E: Результативність",
        )
        self.assertEqual(load_workbook(self.session.path)["Дані_підрозділів"]["E4"].value, 0.89)

    def test_text_percentage_is_copied_verbatim_to_destination_e(self):
        self.set_alias()
        result = self.session.open_input(self.make_source(text_percent=True))
        self.assertEqual(len(result.copied), 1)
        self.assertFalse(result.alias_issues)
        self.assertEqual(self.session.sheet["E4"].value, "89%")

    def test_two_valid_aliases_to_same_destination_are_both_marked_skipped(self):
        self.set_alias()
        self.set_alias(name="Конкуруючий параметр")
        path = self.make_source(second=(18, "Конкуруючий параметр", 0.66, None))
        result = self.session.open_input(path)
        self.assertEqual(result.copied, [])
        self.assertIn(("висновки", 14), result.alias_issues)
        self.assertIn(("висновки", 18), result.alias_issues)
        self.assertIn("Кілька", result.alias_issues[("висновки", 14)])
        self.assertTrue(any("B14" in warning for warning in result.warnings))
        self.assertIsNone(self.session.sheet["E4"].value)
        self.assertEqual(
            self.session.source_alias_targets()[("висновки", 14)],
            "E: Результативність",
        )

    def test_duplicate_name_with_empty_aliased_column_does_not_block_valid_copy(self):
        self.set_alias()
        path = self.make_source(second=(20, PARAMETER, None, "текст без значення в C"))
        result = self.session.open_input(path)
        self.assertEqual(len(result.copied), 1)
        self.assertEqual(result.copied[0].source.row, 14)
        self.assertEqual(self.session.sheet["E4"].value, 0.89)
        self.assertIn(("висновки", 20), result.alias_issues)
        self.assertIn("колонці C", result.alias_issues[("висновки", 20)])
        self.assertNotIn(("висновки", 14), result.alias_issues)

    def test_two_valid_duplicate_names_remain_unresolved(self):
        self.set_alias()
        path = self.make_source(second=(20, PARAMETER, 0.55, "резерв"))
        result = self.session.open_input(path)
        self.assertFalse(result.copied)
        self.assertIn("Декілька", result.alias_issues[("висновки", 14)])
        self.assertIn("Декілька", result.alias_issues[("висновки", 20)])
        self.assertIsNone(self.session.sheet["E4"].value)

    def test_empty_aliased_cell_is_explained_even_with_other_populated_cells(self):
        self.set_alias(column=7)
        result = self.session.open_input(self.make_source())
        self.assertFalse(result.copied)
        self.assertIn("колонці G", result.alias_issues[("висновки", 14)])
        self.assertTrue(any("B14" in warning for warning in result.warnings))
        self.assertIsNone(self.session.sheet["E4"].value)


if __name__ == "__main__":
    unittest.main()
