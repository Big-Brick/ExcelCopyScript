"""Real Tk right-click regression tests for alias deletion on creation tab."""

from __future__ import annotations

import tempfile
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.gui import ExcelCopyApp
from excel_copy_script.models import SourceField
from excel_copy_script.workbook import WorkbookSession


class AliasContextMenuTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.folder = Path(tmp.name)
        self.destination_file = self.folder / "summary.xlsx"
        output = Workbook()
        ws = output.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"], ws["D3"] = (
            "№", "Особовий склад", "Техніка", "Резерв",
        )
        output.save(self.destination_file)

        self.aliases = AliasStore(self.folder / "aliases.json")
        self.session = WorkbookSession(self.destination_file, self.aliases)
        self.aliases.set(
            self.session.template_id, SourceField(1, "Бронетехніка", 42),
            self.session.fields[1],
        )
        self.aliases.set(
            self.session.template_id,
            SourceField(2, "Додаткова назва", 40, sheet="Склад"),
            self.session.fields[1],
        )
        self.aliases.set(
            self.session.template_id,
            SourceField(3, "Підсилення", 21, sheet="Склад"),
            self.session.fields[2],
        )
        source_file = self.folder / "source.xlsx"
        book = Workbook()
        main = book.active
        main.title = "ВИСНОВКИ"
        main["B1"], main["C1"] = "Бронетехніка", 42
        main["B2"], main["C2"] = "Без аліаса", 33
        stock = book.create_sheet("Склад")
        stock["B3"], stock["C3"] = "Підсилення", 21
        book.save(source_file)
        self.result = self.session.open_input(source_file)

        self.app = ExcelCopyApp()
        self.app.geometry("1500x850")
        self.addCleanup(self.app.destroy)
        self.app.session = self.session
        self.app.aliases = self.aliases
        self.app.notebook.select(1)
        self.app._set_source_sheets(self.result.sheets)
        self.app._refresh_tables()
        self.app.update()

    def right_click(self, widget, row_id, column):
        """Deliver a real Button-3 event to the specific row, not selection."""
        bbox = widget.bbox(row_id, column)
        self.assertTrue(bbox, f"Cannot see {row_id} / {column}")
        x, y, width, height = bbox
        with patch.object(tk.Menu, "tk_popup"), patch.object(tk.Menu, "grab_release"):
            widget.event_generate(
                "<Button-3>", x=x + min(8, width // 2), y=y + height // 2,
            )
            self.app.update()
        return getattr(self.app, "_alias_context_menu", None)

    def test_right_click_frozen_source_row_deletes_only_its_alias(self):
        tree = self.app.create_sources.tree
        # A click on the first row must not delete the currently selected second row.
        tree.selection_set("row-2")
        self.app.update()
        menu = self.right_click(tree, "row-1", "name")
        self.assertIsNotNone(menu)
        self.assertEqual(menu.entrycget(0, "label"), "Видалити аліас")
        with patch("excel_copy_script.gui.messagebox.askyesno", return_value=True):
            menu.invoke(0)
        self.assertIsNone(self.aliases.get(self.session.template_id, "Бронетехніка"))
        self.assertIsNotNone(self.aliases.get(self.session.template_id, "Додаткова назва", "Склад"))
        self.assertEqual(self.session.sheet["C4"].value, 42)
        self.assertEqual(tuple(self.app.create_sources.tree.item("row-1", "tags")), ())
        self.assertEqual(self.app.copied_tree.item("0", "values")[-1], "Аліас видалено")

    def test_right_click_values_region_uses_same_source_alias(self):
        self.app.create_sources.selected_sheet.set("Склад")
        self.app.create_sources._on_sheet_changed()
        self.app.update()
        menu = self.right_click(self.app.create_sources.values_tree, "row-3", "C")
        self.assertIsNotNone(menu)
        with patch("excel_copy_script.gui.messagebox.askyesno", return_value=True):
            menu.invoke(0)
        self.assertIsNone(self.aliases.get(self.session.template_id, "Підсилення", "Склад"))
        self.assertIsNotNone(self.aliases.get(self.session.template_id, "Бронетехніка"))
        self.assertEqual(self.session.sheet["D4"].value, 21)
        self.assertEqual(tuple(self.app.create_sources.values_tree.item("row-3", "tags")), ())

    def test_right_click_unaliased_row_has_no_delete_menu(self):
        menu = self.right_click(self.app.create_sources.tree, "row-2", "name")
        self.assertIsNone(menu)

    def test_right_click_output_column_lists_multiple_aliases(self):
        menu = self.right_click(self.app.destination_tree, "3", "Назва")
        self.assertIsNotNone(menu)
        self.assertEqual(menu.entrycget(0, "label"), "Видалити аліас")
        submenu = self.app.nametowidget(str(menu.entrycget(0, "menu")))
        options = [submenu.entrycget(i, "label") for i in range(submenu.index("end") + 1)]
        self.assertEqual(len(options), 2)
        dormant = next(i for i, label in enumerate(options) if "Додаткова назва" in label)
        with patch("excel_copy_script.gui.messagebox.askyesno", return_value=True):
            submenu.invoke(dormant)
        self.assertIsNone(self.aliases.get(self.session.template_id, "Додаткова назва", "Склад"))
        self.assertIsNotNone(self.aliases.get(self.session.template_id, "Бронетехніка"))
        self.assertEqual(self.session.sheet["C4"].value, 42)

    def test_cancel_source_deletion_preserves_green_row(self):
        menu = self.right_click(self.app.create_sources.values_tree, "row-1", "C")
        self.assertIsNotNone(menu)
        with patch("excel_copy_script.gui.messagebox.askyesno", return_value=False):
            menu.invoke(0)
        self.assertIsNotNone(self.aliases.get(self.session.template_id, "Бронетехніка"))
        self.assertEqual(
            tuple(self.app.create_sources.tree.item("row-1", "tags")), ("aliased",),
        )

    def test_right_click_output_without_alias_has_no_menu(self):
        self.assertIsNone(self.right_click(self.app.destination_tree, "2", "Назва"))


if __name__ == "__main__":
    unittest.main()
