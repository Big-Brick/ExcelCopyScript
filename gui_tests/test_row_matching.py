"""Tk integration for optional clearing and ambiguous row selection."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.gui import ExcelCopyApp
from excel_copy_script.models import RowMatch, RowSelectionRequest, SourceField
from excel_copy_script.row_selection_dialog import RowSelectionDialog
from excel_copy_script.workbook import RowSelectionCancelled


class ChooseOutputRowGuiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.summary = self.folder / "summary.xlsx"
        book = Workbook()
        ws = book.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"] = "№", "Перший", "Другий"
        ws["A4"], ws["B4"], ws["C4"] = 1, 10, 999
        ws["A5"], ws["B5"], ws["C5"] = 2, 10, 888
        book.save(self.summary)

        self.input = self.folder / "input.xlsx"
        source = Workbook()
        ws = source.active
        ws.title = "ВИСНОВКИ"
        ws["B1"], ws["C1"] = "Перший показник", 10
        ws["B2"], ws["C2"] = "Другий показник", 20
        source.save(self.input)

        self.aliases = AliasStore(self.folder / "aliases.json")
        self.app = ExcelCopyApp()
        self.addCleanup(self.app.destroy)
        self.app.geometry("1350x830")
        self.app.update()

    def open_summary(self, clear=False):
        self.app.clear_rows_var.set(clear)
        with (
            patch("excel_copy_script.gui.AliasStore.from_profile", return_value=self.aliases),
            patch("excel_copy_script.gui.filedialog.askopenfilename",
                  return_value=str(self.summary)),
        ):
            self.app.open_destination()
        self.aliases.set(
            self.app.session.template_id,
            SourceField(1, "Перший показник", 10),
            self.app.session.fields[0],
        )
        self.aliases.set(
            self.app.session.template_id,
            SourceField(2, "Другий показник", 20),
            self.app.session.fields[1],
        )
        self.app.update()

    def import_file(self):
        with patch("excel_copy_script.gui.filedialog.askopenfilename",
                   return_value=str(self.input)):
            self.app.open_source()
        self.app.update()

    def disk(self, col, row):
        return load_workbook(self.summary)["Дані_підрозділів"][f"{col}{row}"].value

    def test_default_unchecked_preserves_existing_data_on_open(self):
        self.assertFalse(self.app.clear_rows_var.get())
        original = self.summary.read_bytes()
        self.open_summary()
        self.assertEqual(self.summary.read_bytes(), original)
        self.assertIn("Попередні дані збережено", self.app.status_var.get())

    def test_unchecked_partial_match_uses_selection_dialog_callback(self):
        self.open_summary()
        with patch.object(self.app, "_choose_destination_row", return_value=4) as choose:
            self.import_file()
        choose.assert_called_once()
        self.assertEqual(len(choose.call_args.args[0].candidates), 2)
        self.assertTrue(all(not c.complete for c in choose.call_args.args[0].candidates))
        self.assertEqual(self.app.session.current.row, 4)
        self.assertTrue(self.app.session.current.overwrite_existing)
        self.assertEqual(self.disk("C", 4), 20)
        self.assertEqual(self.disk("C", 5), 888)
        self.assertEqual(self.app.opened_sources, [self.input])

    def test_cancelling_dialog_does_not_import_or_remember_input(self):
        self.open_summary()
        original = self.summary.read_bytes()
        previous_status = self.app.status_var.get()
        with patch.object(
            self.app, "_choose_destination_row", side_effect=RowSelectionCancelled,
        ):
            self.import_file()
        self.assertEqual(self.summary.read_bytes(), original)
        self.assertEqual(self.app.opened_sources, [])
        self.assertIsNone(self.app.session.current)
        self.assertEqual(self.app.status_var.get(), previous_status)

    def test_checked_mode_cleans_and_imports_into_first_empty_row(self):
        self.open_summary(clear=True)
        self.assertIsNone(self.disk("B", 4))
        self.assertIsNone(self.disk("C", 5))
        with patch.object(
            self.app, "_choose_destination_row",
            side_effect=AssertionError("Clearing mode must not show chooser"),
        ):
            self.import_file()
        self.assertEqual(self.app.session.current.row, 4)
        self.assertEqual(self.disk("B", 4), 10)
        self.assertEqual(self.disk("C", 4), 20)

    def test_dialog_accepts_overwrite_new_and_cancel(self):
        request = RowSelectionRequest(
            "example.xlsx",
            (RowMatch(4, ("B",), ("C: 999 → 20",), 2),
             RowMatch(5, ("B",), ("C: 888 → 20",), 2)),
        )
        dialog = RowSelectionDialog(self.app, request)
        dialog.tree.selection_set(dialog.tree.get_children()[1])
        self.app.update()
        dialog._overwrite()
        self.assertTrue(dialog.accepted)
        self.assertEqual(dialog.result, 5)

        dialog = RowSelectionDialog(self.app, request)
        self.app.update()
        dialog._new_row()
        self.assertTrue(dialog.accepted)
        self.assertIsNone(dialog.result)

        dialog = RowSelectionDialog(self.app, request)
        self.app.update()
        dialog.destroy()
        self.assertFalse(dialog.accepted)


if __name__ == "__main__":
    unittest.main()
