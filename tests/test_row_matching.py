"""Matching existing destination rows against values resolved by aliases."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.models import SourceField
from excel_copy_script.workbook import RowSelectionCancelled, WorkbookSession


class ExistingRowMatchingTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.folder = Path(tmp.name)
        self.summary = self.folder / "summary.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "№", "Перший", "Другий", "Ручний"
        ws["A4"], ws["A5"], ws["A6"] = 1, 2, 3
        wb.save(self.summary)
        self.aliases = AliasStore(self.folder / "aliases.json")
        self.session = WorkbookSession(self.summary, self.aliases)
        self.aliases.set(
            self.session.template_id,
            SourceField(1, "Параметр B", 10),
            self.session.fields[0],
        )
        self.aliases.set(
            self.session.template_id,
            SourceField(2, "Параметр C", 20),
            self.session.fields[1],
        )
        self.input = self.folder / "source.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "ВИСНОВКИ"
        ws["B1"], ws["C1"] = "Параметр B", 10
        ws["B2"], ws["C2"] = "Параметр C", 20
        ws["B3"], ws["C3"] = "Ручний параметр", 77
        wb.save(self.input)

    def set_rows(self, values):
        # Reopen a fresh session with a realistic pre-existing destination.
        book = load_workbook(self.summary)
        ws = book["Дані_підрозділів"]
        for cell, value in values.items():
            ws[cell] = value
        book.save(self.summary)
        self.session = WorkbookSession(self.summary, self.aliases)

    def disk(self, address):
        return load_workbook(self.summary)["Дані_підрозділів"][address].value

    def test_unique_complete_match_reuses_row_even_with_other_partial_match(self):
        self.set_rows({
            "B4": 10, "C4": 20, "D4": "залишити",
            "B5": 10, "C5": 999,
        })
        called = []
        def selector(request):
            called.append(request)
            raise AssertionError("Unique full match should never prompt")
        result = self.session.open_input(self.input, row_selector=selector)
        self.assertEqual(called, [])
        self.assertEqual(result.row, 4)
        self.assertTrue(result.overwrite_existing)
        self.assertEqual(self.disk("D4"), "залишити")
        self.assertEqual(self.disk("C5"), 999)
        self.assertIsNone(self.disk("B6"))

    def test_multiple_full_matches_show_dialog_and_allow_specific_overwrite(self):
        self.set_rows({"B4": 10, "C4": 20, "B5": 10, "C5": 20})
        seen = []
        def selector(request):
            seen.append(request)
            return 5
        result = self.session.open_input(self.input, row_selector=selector)
        self.assertEqual(result.row, 5)
        self.assertEqual([row.row for row in seen[0].candidates], [4, 5])
        self.assertTrue(all(row.complete for row in seen[0].candidates))
        self.assertTrue(result.overwrite_existing)
        self.assertIsNone(self.disk("B6"))

    def test_multiple_full_matches_user_can_choose_new_row(self):
        self.set_rows({"B4": 10, "C4": 20, "B5": 10, "C5": 20})
        result = self.session.open_input(self.input, row_selector=lambda req: None)
        self.assertEqual(result.row, 6)
        self.assertFalse(result.overwrite_existing)
        self.assertEqual(self.disk("B6"), 10)
        self.assertEqual(self.disk("C6"), 20)

    def test_partial_match_prompts_and_overwrites_selected_values_only(self):
        self.set_rows({
            "B4": 10, "C4": 99, "D4": "keep",
            "B5": 10, "C5": 88, "D5": "unchanged",
        })
        seen = []
        def selector(req):
            seen.append(req)
            return 4
        result = self.session.open_input(self.input, row_selector=selector)
        self.assertEqual(result.row, 4)
        self.assertTrue(result.overwrite_existing)
        self.assertEqual(len(seen[0].candidates), 2)
        self.assertEqual(seen[0].candidates[0].matching, ("B",))
        self.assertTrue(any("C:" in text for text in seen[0].candidates[0].differing))
        self.assertEqual(self.disk("C4"), 20)
        self.assertEqual(self.disk("D4"), "keep")
        self.assertEqual(self.disk("C5"), 88)

    def test_partial_match_user_can_choose_new(self):
        self.set_rows({"B4": 10, "C4": 99})
        result = self.session.open_input(self.input, row_selector=lambda request: None)
        self.assertEqual(result.row, 5)
        self.assertEqual(self.disk("C4"), 99)
        self.assertEqual(self.disk("C5"), 20)

    def test_no_match_writes_new_without_prompt(self):
        self.set_rows({"B4": 44, "C4": 99})
        result = self.session.open_input(
            self.input,
            row_selector=lambda request: self.fail("No matches: no dialog"),
        )
        self.assertEqual(result.row, 5)
        self.assertFalse(result.overwrite_existing)
        self.assertEqual(self.disk("B4"), 44)
        self.assertEqual(self.disk("B5"), 10)

    def test_ambiguous_matching_never_writes_without_choice(self):
        self.set_rows({"B4": 10, "C4": 99})
        before = self.summary.read_bytes()
        with self.assertRaisesRegex(ValueError, "Потрібно вибрати"):
            self.session.open_input(self.input)
        self.assertEqual(self.summary.read_bytes(), before)
        self.assertIsNone(self.session.current)

    def test_cancellation_preserves_disk_and_previous_session(self):
        self.set_rows({"B4": 10, "C4": 99})
        before = self.summary.read_bytes()
        with self.assertRaises(RowSelectionCancelled):
            self.session.open_input(
                self.input,
                row_selector=lambda request: (_ for _ in ()).throw(RowSelectionCancelled()),
            )
        self.assertEqual(self.summary.read_bytes(), before)
        self.assertIsNone(self.session.current)

    def test_unlisted_row_selection_is_rejected_before_any_write(self):
        self.set_rows({"B4": 10, "C4": 99})
        before = self.summary.read_bytes()
        with self.assertRaisesRegex(ValueError, "недопустимий"):
            self.session.open_input(self.input, row_selector=lambda request: 17)
        self.assertEqual(self.summary.read_bytes(), before)

    def test_manual_copy_overwrites_chosen_existing_row_and_keeps_other_columns(self):
        self.set_rows({"B4": 10, "C4": 99, "D4": "замінити"})
        result = self.session.open_input(self.input, row_selector=lambda request: 4)
        self.assertEqual(result.row, 4)
        self.session.copy_manually(3, 4)
        self.assertEqual(self.disk("D4"), 77)
        self.assertEqual(self.disk("C4"), 20)

    def test_disabled_matching_appends_even_when_full_match_exists(self):
        self.set_rows({"B4": 10, "C4": 20})
        result = self.session.open_input(
            self.input, match_existing=False,
            row_selector=lambda request: self.fail("No dialog in clearing mode"),
        )
        self.assertEqual(result.row, 5)
        self.assertEqual(self.disk("B4"), 10)
        self.assertEqual(self.disk("B5"), 10)

    def test_row_comparison_ignores_unaliased_columns(self):
        self.set_rows({"B4": 10, "C4": 20, "D4": "arbitrary"})
        result = self.session.open_input(self.input)
        self.assertEqual(result.row, 4)
        self.assertEqual(self.disk("D4"), "arbitrary")

    def test_repeat_restart_without_cleaning_merges_instead_of_duplicating(self):
        self.session.open_input(self.input)
        rebuilt = self.session.restart_with_inputs([self.input])
        self.assertEqual(rebuilt.current.row, 4)
        self.assertEqual(rebuilt.next_empty_row(), 5)
        self.assertEqual(self.disk("B4"), 10)

    def test_restart_rollback_if_row_choice_cancelled(self):
        self.set_rows({"B4": 10, "C4": 99})
        before = self.summary.read_bytes()
        with self.assertRaises(RowSelectionCancelled):
            self.session.restart_with_inputs(
                [self.input],
                row_selector=lambda req: (_ for _ in ()).throw(RowSelectionCancelled()),
            )
        self.assertEqual(self.summary.read_bytes(), before)
        self.assertEqual(list(self.folder.glob(".excel-copy-restart-*.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
