"""Transactional full restart: replay all inputs using current saved aliases."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.models import SourceField
from excel_copy_script.workbook import WorkbookSession


class FullRestartTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.destination = self.folder / "summary.xlsx"
        self.alias_file = self.folder / "aliases.json"
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        ws["A3"], ws["B3"], ws["C3"], ws["D3"] = (
            "№", "Особовий склад", "Техніка", "Резерв",
        )
        ws["A4"], ws["A5"] = 1, 2
        wb.save(self.destination)
        self.aliases = AliasStore(self.alias_file)
        self.session = WorkbookSession(self.destination, self.aliases)
        self.source_a = self.folder / "a.xlsx"
        self.source_b = self.folder / "b.xlsx"
        self.make_source(self.source_a, "ВИСНОВКИ", [
            (1, "Чисельність", 11), (2, "Запас", 100),
        ])
        self.make_source(self.source_b, "Склад", [(1, "Бронетехніка", 22)])
        self.aliases.set(
            self.session.template_id,
            SourceField(1, "Чисельність", 11),
            self.session.fields[0],
        )
        self.aliases.set(
            self.session.template_id,
            SourceField(1, "Бронетехніка", 22, sheet="Склад"),
            self.session.fields[1],
        )

    @staticmethod
    def make_source(path, sheet_name, entries):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row, name, value in entries:
            ws.cell(row, 2).value = name
            ws.cell(row, 3).value = value
        wb.save(path)

    def values(self):
        ws = load_workbook(self.destination)["Дані_підрозділів"]
        return {address: ws[address].value for address in (
            "A4", "A5", "B4", "C4", "D4", "B5", "C5", "D5",
        )}

    def assert_no_stage_files(self):
        self.assertFalse(list(self.folder.glob(".excel-copy-restart-*.xlsx")))
        self.assertFalse(list(self.folder.glob("*.backup-*.xlsx")))

    def test_replays_sources_in_order_with_newly_added_and_deleted_aliases(self):
        self.session.open_input(self.source_a)
        self.session.open_input(self.source_b)
        self.assertEqual(self.values()["B4"], 11)
        self.assertEqual(self.values()["C5"], 22)
        self.assertTrue(self.aliases.delete(self.session.template_id, "Чисельність"))
        self.aliases.set(
            self.session.template_id,
            SourceField(2, "Запас", 100),
            self.session.fields[2],
        )
        original = self.session
        rebuilt = original.restart_with_inputs([self.source_a, self.source_b])
        self.assertIsNot(rebuilt, original)
        self.assertIsNot(rebuilt.aliases, self.aliases)
        self.assertEqual(rebuilt.path, self.destination)
        self.assertEqual(rebuilt.current.filename, "b.xlsx")
        self.assertEqual(rebuilt.current.row, 5)
        data = self.values()
        self.assertEqual((data["A4"], data["A5"]), (1, 2))
        self.assertIsNone(data["B4"])
        self.assertEqual(data["D4"], 100)
        self.assertEqual(data["C5"], 22)
        self.assertEqual(rebuilt.next_empty_row(), 6)
        self.assert_no_stage_files()

    def test_restart_picks_up_modified_input_file_on_disk(self):
        self.session.open_input(self.source_a)
        self.make_source(self.source_a, "ВИСНОВКИ", [
            (1, "Чисельність", 77), (2, "Запас", 100),
        ])
        rebuilt = self.session.restart_with_inputs([self.source_a])
        self.assertEqual(self.values()["B4"], 77)
        self.assertEqual(rebuilt.current.row, 4)
        self.assert_no_stage_files()

    def test_restart_without_inputs_only_clears_previous_rows(self):
        self.session.open_input(self.source_a)
        result = self.session.restart_with_inputs([])
        self.assertIsNone(result.current)
        self.assertIsNone(self.values()["B4"])
        self.assertEqual(self.values()["A4"], 1)
        self.assertEqual(result.next_empty_row(), 4)
        self.assert_no_stage_files()

    def test_repeated_source_file_path_is_replayed_twice(self):
        rebuilt = self.session.restart_with_inputs([self.source_a, self.source_a])
        self.assertEqual(self.values()["B4"], 11)
        self.assertEqual(self.values()["B5"], 11)
        self.assertEqual(rebuilt.current.row, 5)

    def test_missing_source_does_not_modify_destination(self):
        self.session.open_input(self.source_a)
        before = self.destination.read_bytes()
        original_current = self.session.current
        with self.assertRaisesRegex(ValueError, "missing.xlsx"):
            self.session.restart_with_inputs([
                self.source_a, self.folder / "missing.xlsx",
            ])
        self.assertEqual(self.destination.read_bytes(), before)
        self.assertIs(self.session.current, original_current)
        self.assert_no_stage_files()

    def test_corrupt_source_does_not_commit_partial_import(self):
        self.session.open_input(self.source_a)
        original = self.destination.read_bytes()
        corrupt = self.folder / "bad.xlsx"
        corrupt.write_text("not an Excel workbook", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "bad.xlsx"):
            self.session.restart_with_inputs([self.source_a, corrupt])
        self.assertEqual(self.destination.read_bytes(), original)
        self.assert_no_stage_files()


if __name__ == "__main__":
    unittest.main()
