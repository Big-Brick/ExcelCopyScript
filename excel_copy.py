"""Interactive transfer of values from ВИСНОВКИ into Дані_підрозділів."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import tkinter as tk
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DESTINATION_SHEET = "Дані_підрозділів"
SOURCE_SHEET = "ВИСНОВКИ"
HEADER_ROW = 3
FIRST_DATA_ROW = 4
FILE_TYPES = [("Книги Excel", "*.xlsx")]


def normalize(value: str) -> str:
    """Ignore casing, Unicode presentation variants and repeated whitespace."""
    return " ".join(unicodedata.normalize("NFKC", value).casefold().split())


def present(value: object) -> str:
    return "" if value is None else str(value).replace("\n", " ")


def nonempty(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def text_as_excel_value(value: object) -> object:
    """Never turn text copied from a source cell into a new Excel formula."""
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


@dataclass(frozen=True)
class DestinationField:
    column: int
    name: str

    @property
    def key(self) -> str:
        return normalize(self.name)

    @property
    def letter(self) -> str:
        return get_column_letter(self.column)


@dataclass(frozen=True)
class SourceField:
    row: int
    name: str
    value: object

    @property
    def key(self) -> str:
        return normalize(self.name)


@dataclass(frozen=True)
class CopyRecord:
    source: SourceField
    destination: DestinationField
    method: str


@dataclass
class ImportResult:
    filename: str
    row: int
    copied: list[CopyRecord] = field(default_factory=list)
    unmatched_destinations: list[DestinationField] = field(default_factory=list)
    unmatched_sources: list[SourceField] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class AliasStore:
    """Persistent input-name -> destination-header mapping, scoped to a header set."""

    def __init__(self, path: Path | None = None):
        self.path = path or (Path.home() / ".config" / "excel-copy-script" / "aliases.json")
        if self.path.exists():
            try:
                data = json.loads(self.path.read_text(encoding="utf-8"))
            except (OSError, ValueError) as exc:
                raise ValueError(f"Не вдалося прочитати файл аліасів {self.path}: {exc}") from exc
            if not isinstance(data, dict) or data.get("version") != 1 or not isinstance(data.get("templates"), dict):
                raise ValueError(f"Непідтримуваний формат файлу аліасів: {self.path}")
            self.data = data
        else:
            self.data = {"version": 1, "templates": {}}

    @staticmethod
    def template_id(fields: list[DestinationField]) -> str:
        headers = sorted(field.key for field in fields)
        return hashlib.sha256(json.dumps(headers, ensure_ascii=False).encode("utf-8")).hexdigest()

    def get(self, template_id: str, source_name: str) -> dict | None:
        entry = self.data["templates"].get(template_id, {}).get(normalize(source_name))
        return entry if isinstance(entry, dict) else None

    def set(self, template_id: str, source: SourceField, destination: DestinationField) -> None:
        key = source.key
        template = self.data["templates"].setdefault(template_id, {})
        previous = template.get(key)
        template[key] = {"header": destination.key, "column": destination.column}
        try:
            self._save()
        except Exception:
            if previous is None:
                del template[key]
            else:
                template[key] = previous
            raise

    def _save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=self.path.parent,
                                             prefix=".aliases-", suffix=".tmp", delete=False) as stream:
                tmp_path = Path(stream.name)
                json.dump(self.data, stream, ensure_ascii=False, indent=2)
                stream.write("\n")
            os.replace(tmp_path, self.path)
        finally:
            if tmp_path is not None:
                tmp_path.unlink(missing_ok=True)


class WorkbookSession:
    def __init__(self, path: Path, aliases: AliasStore):
        self.path = Path(path)
        if self.path.suffix.lower() != ".xlsx":
            raise ValueError("Підтримуються лише файли .xlsx")
        self.workbook = load_workbook(self.path)
        if DESTINATION_SHEET not in self.workbook:
            raise ValueError(f"У підсумковому файлі немає аркуша «{DESTINATION_SHEET}»")
        self.sheet = self.workbook[DESTINATION_SHEET]
        self.fields = [DestinationField(cell.column, cell.value.strip())
                       for cell in self.sheet[HEADER_ROW]
                       if isinstance(cell.value, str) and cell.value.strip()]
        if not self.fields:
            raise ValueError(f"У рядку {HEADER_ROW} немає текстових заголовків")
        self.aliases = aliases
        self.template_id = aliases.template_id(self.fields)
        self.current: ImportResult | None = None
        self.backup_path: Path | None = None

    def next_empty_row(self) -> int:
        # Consider all cells, not only mapped columns: never overwrite a partially populated row.
        for row in range(FIRST_DATA_ROW, self.sheet.max_row + 1):
            if all(not nonempty(self.sheet.cell(row, col).value)
                   for col in range(1, self.sheet.max_column + 1)):
                return row
        return max(FIRST_DATA_ROW, self.sheet.max_row + 1)

    def _read_source(self, path: Path) -> tuple[list[SourceField], list[str]]:
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Підтримуються лише файли .xlsx")
        if path.resolve() == self.path.resolve():
            raise ValueError("Вхідний і підсумковий файли не можуть збігатися")
        cached = load_workbook(path, read_only=True, data_only=True)
        formulas = None
        try:
            if SOURCE_SHEET not in cached:
                raise ValueError(f"У файлі {path.name} немає аркуша «{SOURCE_SHEET}»")
            # Read formula metadata separately: openpyxl does not calculate formulas.
            formulas = load_workbook(path, read_only=True, data_only=False)
            cached_ws = cached[SOURCE_SHEET]
            formula_ws = formulas[SOURCE_SHEET]
            entries: list[SourceField] = []
            warnings: list[str] = []
            for row, (name, cached_value) in enumerate(
                cached_ws.iter_rows(min_col=2, max_col=3, values_only=True), start=1
            ):
                if not isinstance(name, str) or not name.strip():
                    continue
                if not nonempty(cached_value):
                    formula_cell = formula_ws.cell(row, 3)
                    if formula_cell.data_type == "f":
                        warnings.append(f"C{row}: формула без збереженого результату — пропущено")
                    continue
                entries.append(SourceField(row, name.strip(), cached_value))
            return entries, warnings
        finally:
            cached.close()
            if formulas is not None:
                formulas.close()

    def _destination_for(self, source: SourceField) -> tuple[DestinationField | None, str]:
        alias = self.aliases.get(self.template_id, source.name)
        if alias:
            candidates = [field for field in self.fields if field.key == alias.get("header")]
            if len(candidates) == 1:
                return candidates[0], "Аліас"
            if len(candidates) > 1:
                pinned = [field for field in candidates if field.column == alias.get("column")]
                if len(pinned) == 1:
                    return pinned[0], "Аліас"
        exact = [field for field in self.fields if field.key == source.key]
        if len(exact) == 1:
            return exact[0], "Збіг назви"
        return None, ""

    def _write(self, row: int, records: list[CopyRecord]) -> None:
        if not records:
            return
        before = {}
        for record in records:
            column = record.destination.column
            cell = self.sheet.cell(row, column)
            if nonempty(cell.value):
                raise ValueError(f"Комірка {cell.coordinate} вже заповнена")
            before[column] = cell.value
            cell.value = text_as_excel_value(record.source.value)
        try:
            self._save()
        except Exception:
            for column, old_value in before.items():
                self.sheet.cell(row, column).value = old_value
            raise

    def _save(self) -> None:
        if self.backup_path is None:
            suffix = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
            backup = self.path.with_name(f"{self.path.stem}.backup-{suffix}.xlsx")
            shutil.copy2(self.path, backup)
            self.backup_path = backup
        tmp_path = None
        try:
            with tempfile.NamedTemporaryFile(dir=self.path.parent, prefix=".excel-copy-",
                                             suffix=".xlsx", delete=False) as file:
                tmp_path = Path(file.name)
            self.workbook.save(tmp_path)
            os.chmod(tmp_path, self.path.stat().st_mode)
            os.replace(tmp_path, self.path)
        finally:
            if tmp_path is not None:
                tmp_path.unlink(missing_ok=True)

    def open_input(self, path: Path) -> ImportResult:
        source_fields, warnings = self._read_source(Path(path))
        target_row = self.next_empty_row()
        duplicates = Counter(field.key for field in source_fields)
        proposals: dict[int, list[CopyRecord]] = defaultdict(list)
        for source in source_fields:
            if duplicates[source.key] > 1:
                continue
            destination, method = self._destination_for(source)
            if destination is not None:
                proposals[destination.column].append(CopyRecord(source, destination, method))
        safe_matches = [matches[0] for matches in proposals.values() if len(matches) == 1]
        safe_matches.sort(key=lambda rec: rec.destination.column)
        self._write(target_row, safe_matches)
        result = ImportResult(filename=Path(path).name, row=target_row, copied=safe_matches,
                              warnings=warnings)
        if any(count > 1 for count in duplicates.values()):
            result.warnings.append("Повторювані назви у вхідному файлі не зіставлено автоматично")
        if any(len(matches) > 1 for matches in proposals.values()):
            result.warnings.append("Кілька вхідних параметрів претендують на один стовпчик — пропущено")
        used_rows = {record.source.row for record in safe_matches}
        used_columns = {record.destination.column for record in safe_matches}
        result.unmatched_sources = [source for source in source_fields if source.row not in used_rows]
        result.unmatched_destinations = [dest for dest in self.fields if dest.column not in used_columns]
        self.current = result
        return result

    def copy_manually(self, source_row: int, destination_col: int) -> str | None:
        result = self.current
        if result is None:
            raise ValueError("Спочатку відкрийте вхідний файл")
        source = next((s for s in result.unmatched_sources if s.row == source_row), None)
        destination = next((d for d in result.unmatched_destinations if d.column == destination_col), None)
        if source is None or destination is None:
            raise ValueError("Вибраний параметр уже був скопійований")
        record = CopyRecord(source, destination, "Вручну")
        self._write(result.row, [record])
        result.copied.append(record)
        result.unmatched_sources.remove(source)
        result.unmatched_destinations.remove(destination)
        try:
            self.aliases.set(self.template_id, source, destination)
        except (OSError, ValueError) as exc:
            return f"Значення скопійовано, але аліас не збережено: {exc}"
        return None


class ExcelCopyApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Перенесення даних Excel")
        self.geometry("1260x720")
        self.minsize(900, 510)
        self.session: WorkbookSession | None = None
        self.aliases: AliasStore | None = None
        self.destination_var = tk.StringVar(value="Підсумковий файл не вибрано")
        self.source_var = tk.StringVar(value="Вхідний файл не вибрано")
        self.status_var = tk.StringVar(value="Виберіть підсумковий файл, потім вхідний.")
        self._build_ui()

    def _build_ui(self) -> None:
        top = ttk.Frame(self, padding=12)
        top.pack(fill="x")
        ttk.Button(top, text="1. Відкрити підсумковий файл", command=self.open_destination).pack(side="left")
        self.source_button = ttk.Button(top, text="2. Відкрити вхідний / наступний файл",
                                        command=self.open_source, state="disabled")
        self.source_button.pack(side="left", padx=(12, 0))
        ttk.Label(self, textvariable=self.destination_var, padding=(12, 0)).pack(anchor="w")
        ttk.Label(self, textvariable=self.source_var, padding=(12, 2)).pack(anchor="w")

        panes = ttk.Panedwindow(self, orient="horizontal")
        panes.pack(fill="both", expand=True, padx=12, pady=(8, 4))
        copied_frame, self.copied_tree = self._make_table(
            panes, "Скопійовані значення", ("Звідки", "Куди", "Значення", "Спосіб"), (180, 190, 140, 130))
        destination_frame, self.destination_tree = self._make_table(
            panes, "Без відповідника: вихідні стовпчики", ("Колонка", "Назва"), (80, 260))
        source_frame, self.source_tree = self._make_table(
            panes, "Без відповідника: вхідні рядки", ("Рядок", "Назва", "Значення"), (70, 240, 110))
        for frame in (copied_frame, destination_frame, source_frame):
            panes.add(frame, weight=1)

        bottom = ttk.Frame(self, padding=(12, 8, 12, 12))
        bottom.pack(fill="x")
        self.copy_button = ttk.Button(bottom, text="Копіювати вибрану пару + запам'ятати аліас",
                                      command=self.copy_selected, state="disabled")
        self.copy_button.pack(anchor="w")
        ttk.Label(bottom, text="Виберіть один рядок у кожному зі списків «Без відповідника».").pack(anchor="w", pady=(5, 3))
        ttk.Separator(bottom).pack(fill="x", pady=(3, 8))
        ttk.Label(bottom, textvariable=self.status_var, wraplength=1150).pack(anchor="w")
        self.destination_tree.bind("<<TreeviewSelect>>", self._on_selection)
        self.source_tree.bind("<<TreeviewSelect>>", self._on_selection)

    def _make_table(self, parent, title, headings, widths):
        frame = ttk.Frame(parent, padding=4)
        ttk.Label(frame, text=title, font=("", 10, "bold")).pack(anchor="w", pady=(0, 5))
        container = ttk.Frame(frame)
        container.pack(fill="both", expand=True)
        tree = ttk.Treeview(container, columns=headings, show="headings", selectmode="browse")
        for heading, width in zip(headings, widths):
            tree.heading(heading, text=heading)
            tree.column(heading, width=width, minwidth=60, stretch=True)
        scrollbar_y = ttk.Scrollbar(container, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar_y.set)
        tree.pack(side="left", fill="both", expand=True)
        scrollbar_y.pack(side="right", fill="y")
        return frame, tree

    def _on_selection(self, _event=None):
        can_copy = bool(self.session and self.session.current and
                        self.destination_tree.selection() and self.source_tree.selection())
        self.copy_button.configure(state="normal" if can_copy else "disabled")

    def open_destination(self):
        chosen = filedialog.askopenfilename(title="Підсумковий Excel-файл", filetypes=FILE_TYPES)
        if not chosen:
            return
        try:
            aliases = AliasStore()
            session = WorkbookSession(Path(chosen), aliases)
        except Exception as exc:
            messagebox.showerror("Помилка відкриття", str(exc))
            return
        self.aliases, self.session = aliases, session
        self.destination_var.set(f"Підсумковий: {session.path}")
        self.source_var.set("Вхідний файл не вибрано")
        self.source_button.configure(state="normal")
        self._clear_tables()
        self.status_var.set(f"Знайдено {len(session.fields)} стовпчиків. Наступний рядок: {session.next_empty_row()}.")

    def open_source(self):
        if self.session is None:
            return
        chosen = filedialog.askopenfilename(title="Вхідний Excel-файл", filetypes=FILE_TYPES,
                                            initialdir=str(self.session.path.parent))
        if not chosen:
            return
        try:
            result = self.session.open_input(Path(chosen))
        except Exception as exc:
            messagebox.showerror("Помилка імпорту", str(exc))
            return
        self.source_var.set(f"Вхідний: {chosen} → рядок {result.row}")
        self._refresh_tables()

    def _clear_tables(self):
        for tree in (self.copied_tree, self.destination_tree, self.source_tree):
            tree.delete(*tree.get_children())
        self._on_selection()

    def _refresh_tables(self):
        self._clear_tables()
        if self.session is None or self.session.current is None:
            return
        result = self.session.current
        for record in result.copied:
            self.copied_tree.insert("", "end", values=(f"B{record.source.row}: {record.source.name}",
                f"{record.destination.letter}: {record.destination.name}", present(record.source.value), record.method))
        for dest in result.unmatched_destinations:
            self.destination_tree.insert("", "end", iid=str(dest.column), values=(dest.letter, dest.name))
        for source in result.unmatched_sources:
            self.source_tree.insert("", "end", iid=str(source.row), values=(source.row, source.name, present(source.value)))
        status = (f"{result.filename} → рядок {result.row}. Скопійовано: {len(result.copied)}; "
                  f"не зіставлено вихідних стовпчиків: {len(result.unmatched_destinations)}; "
                  f"вхідних рядків: {len(result.unmatched_sources)}. Зміни автоматично збережено.")
        if result.warnings:
            status += " Увага: " + "; ".join(result.warnings[:4])
        self.status_var.set(status)

    def copy_selected(self):
        if self.session is None:
            return
        sources, destinations = self.source_tree.selection(), self.destination_tree.selection()
        if not sources or not destinations:
            return
        try:
            warning = self.session.copy_manually(int(sources[0]), int(destinations[0]))
        except Exception as exc:
            messagebox.showerror("Помилка копіювання", str(exc))
            return
        self._refresh_tables()
        if warning:
            messagebox.showwarning("Аліас не збережено", warning)


def main() -> None:
    ExcelCopyApp().mainloop()


if __name__ == "__main__":
    main()
