import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_copy_script.aliases import AliasStore
from excel_copy_script.utils import normalize
from excel_copy_script.workbook import WorkbookSession


class TransferTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.target = self.directory / "summary.xlsx"
        self.aliases = AliasStore(self.directory / "config" / "aliases.json")
        wb = Workbook()
        ws = wb.active
        ws.title = "Дані_підрозділів"
        ws["A3"] = "Підрозділ"
        ws["B3"] = "Особовий склад"
        ws["C3"] = "Техніка"
        ws["D3"] = "Резерв"
        ws["A4"] = "Вже заповнено"
        ws["B4"] = 99
        wb.save(self.target)
        self.session = WorkbookSession(self.target, self.aliases)

    def source(self, name, rows):
        path = self.directory / name
        wb = Workbook()
        ws = wb.active
        ws.title = "ВИСНОВКИ"
        for row, title, value in rows:
            ws.cell(row, 2, title)
            ws.cell(row, 3, value)
        wb.save(path)
        return path

    def test_exact_manual_alias_next_file_and_reopen(self):
        first = self.source("first.xlsx", [(1, "Особовий  склад", 13), (2, "Бронетехніка", 7),
                                               (4, "Розділ звіту", None)])
        result = self.session.open_input(first)
        self.assertEqual(result.row, 5)
        self.assertEqual(len(result.copied), 1)
        self.assertEqual(result.copied[0].destination.column, 2)
        self.assertEqual({item.row for item in result.unmatched_sources}, {2})
        self.assertIsNone(self.session.copy_manually(2, 3))
        self.assertTrue(self.aliases.path.exists())
        self.assertEqual(len(self.session.current.copied), 2)

        # Persistent aliases are resolved by header, not by source/destination row.
        reloaded_aliases = AliasStore(self.aliases.path)
        self.session = WorkbookSession(self.target, reloaded_aliases)
        second = self.source("second.xlsx", [(6, "особовий склад", 15), (7, "Бронетехніка", 8)])
        result = self.session.open_input(second)
        self.assertEqual(result.row, 6)
        self.assertEqual({item.method for item in result.copied}, {"Збіг назви", "Аліас"})
        ws = load_workbook(self.target, data_only=True)["Дані_підрозділів"]
        self.assertEqual(ws["B4"].value, 99)
        self.assertEqual(ws["B5"].value, 13)
        self.assertEqual(ws["C5"].value, 7)
        self.assertEqual(ws["B6"].value, 15)
        self.assertEqual(ws["C6"].value, 8)
        self.assertEqual(len(list(self.directory.glob("summary.backup-*.xlsx"))), 2)

    def test_duplicate_source_names_are_not_auto_matched(self):
        file = self.source("duplicate.xlsx", [(1, "Техніка", 1), (3, " ТЕХНІКА ", 2)])
        result = self.session.open_input(file)
        self.assertEqual(len(result.copied), 0)
        self.assertEqual(len(result.unmatched_sources), 2)
        self.assertTrue(any("Повторювані" in note for note in result.warnings))
        self.assertFalse(list(self.directory.glob("summary.backup-*.xlsx")))

    def test_first_completely_empty_row_skips_partial(self):
        ws = self.session.sheet
        ws["D5"] = 1
        self.assertEqual(self.session.next_empty_row(), 6)

    def test_aliases_scoped_to_output_template(self):
        first = self.source("match.xlsx", [(1, "Нова назва", 55)])
        self.session.open_input(first)
        self.session.copy_manually(1, 2)
        other = Workbook()
        ws = other.active
        ws.title = "Дані_підрозділів"
        ws["A3"] = "Інша назва колонки"
        different = self.directory / "other.xlsx"
        other.save(different)
        another_session = WorkbookSession(different, AliasStore(self.aliases.path))
        result = another_session.open_input(first)
        self.assertFalse(result.copied)

    def test_conflicting_candidates_do_not_overwrite_each_other(self):
        source = self.source("conflict.xlsx", [(1, "Техніка", 11), (2, "Запас техніки", 22)])
        self.aliases.set(self.session.template_id,
                         type("S", (), {"key": normalize("Запас техніки")})(), self.session.fields[2])
        result = self.session.open_input(source)
        self.assertFalse(result.copied)
        self.assertEqual(len(result.unmatched_sources), 2)
        self.assertTrue(any("претендують" in note for note in result.warnings))

    def test_missing_formula_cache_warns_and_does_not_copy(self):
        file = self.source("formulas.xlsx", [(3, "Техніка", "=2+2")])
        result = self.session.open_input(file)
        self.assertFalse(result.copied)
        self.assertTrue(any("формула" in warning for warning in result.warnings))

    def test_untrusted_string_does_not_become_formula(self):
        source = self.source("literal.xlsx", [(1, "Техніка", "'=cmd")])
        result = self.session.open_input(source)
        self.assertEqual(len(result.copied), 1)
        wb = load_workbook(self.target)
        self.assertNotEqual(wb["Дані_підрозділів"]["C5"].data_type, "f")

    def test_bad_alias_file_fails_loudly(self):
        self.aliases.path.parent.mkdir(parents=True, exist_ok=True)
        self.aliases.path.write_text("not json", encoding="utf-8")
        with self.assertRaises(ValueError):
            AliasStore(self.aliases.path)


if __name__ == "__main__":
    unittest.main()
